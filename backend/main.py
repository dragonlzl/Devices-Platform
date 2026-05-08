import asyncio
import json
import logging
import re
from io import BytesIO
from datetime import datetime, timedelta, timezone
from typing import Any, Dict, List, Optional, Union

import httpx
import os
from pathlib import Path
from fastapi import BackgroundTasks, FastAPI, HTTPException, Query, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook

from .db import db_session, init_db, now_iso
from .llm import LLMError, call_llm
from .notify import (
    DEFAULT_BORROW_ADMIN_URL,
    NotifyError,
    PortalNotificationError,
    send_feishu_message,
    send_portal_notification,
)
from .schemas import (
    BorrowRequest,
    BorrowerChangeRequest,
    DeviceCreate,
    DeviceUpdate,
    ExtendRequest,
    LLMModelCreate,
    LLMModelAssignRequest,
    LLMModelUpdate,
    LLMTestRequest,
    NotificationSettingsUpdate,
    SettingUpdate,
    SystemCreate,
    SystemDeleteRequest,
    SystemUpdate,
    VendorCreate,
    VendorDeleteRequest,
    VendorUpdate,
    VersionCreate,
    VersionDeleteRequest,
    VersionUpdate,
    WebhookNotificationSettingsUpdate,
)


BASE_DIR = Path(__file__).resolve()
BACKEND_DIR = BASE_DIR.parent
REPO_DIR = BACKEND_DIR.parent


def _load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        key, value = stripped.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


_load_env_file(REPO_DIR / ".env")

PORTAL_JWT_VERIFY_URL = os.environ.get(
    "PORTAL_JWT_VERIFY_URL",
    "http://192.168.50.10:8756/api/sso/jwt/verify",
)
DEFAULT_PORTAL_NOTIFICATION_SERVICE_ID = "device-borrow-service"
PORTAL_NOTIFICATION_SERVICE_ID_ENV = "PORTAL_NOTIFICATION_SERVICE_ID"
PORTAL_NOTIFICATION_SERVICE_TOKEN_ENV = "PORTAL_NOTIFICATION_SERVICE_TOKEN"
RESIDENT_DEVICE_STATUS = "被常驻"
logger = logging.getLogger(__name__)

NOTIFICATION_SETTINGS_KEY = "portal_notification_params"
WEBHOOK_NOTIFICATION_SETTINGS_KEY = "feishu_webhook_card_params"
BORROW_ADMIN_URL_SETTING_KEY = "borrow_admin_url"
NOTIFICATION_COLOR_OPTIONS = ("blue", "green", "yellow", "orange", "red", "grey")
NOTIFICATION_PARAM_FIELDS = ("card_title", "status", "card_color", "status_color")
NOTIFICATION_COLOR_FIELDS = {"card_color", "status_color"}
WEBHOOK_NOTIFICATION_PARAM_FIELDS = ("card_title", "body_template", "card_color")
NOTIFICATION_TRIGGER_DEFAULTS: Dict[str, Dict[str, Any]] = {
    "borrow_submit": {
        "label": "待借申请提交",
        "description": "借用页提交待借申请后通知申请人。",
        "params": {
            "card_title": "借用申请已提交",
            "status": "待管理员确认设备状态",
            "card_color": "blue",
            "status_color": "blue",
        },
    },
    "borrow_approve": {
        "label": "借用确认成功",
        "description": "管理员确认借出后通知借用人。",
        "params": {
            "card_title": "借用成功",
            "status": "借用成功，请联系 @林镇龙 领取设备",
            "card_color": "green",
            "status_color": "green",
        },
    },
    "borrow_cancel": {
        "label": "借用申请取消",
        "description": "管理员取消待借申请后通知申请人。",
        "params": {
            "card_title": "借用失败",
            "status": "借用失败，目标设备状态异常，请选另一台设备",
            "card_color": "red",
            "status_color": "red",
        },
    },
    "extend": {
        "label": "延期成功",
        "description": "借用人延期归还时间后通知借用人。",
        "params": {
            "card_title": "归还时间延长成功",
            "status": "归还时间延长成功",
            "card_color": "green",
            "status_color": "green",
        },
    },
    "return": {
        "label": "归还成功",
        "description": "管理员归还设备后通知原借用人。",
        "params": {
            "card_title": "归还成功",
            "status": "归还成功",
            "card_color": "green",
            "status_color": "green",
        },
    },
    "overdue": {
        "label": "逾期提醒",
        "description": "后台检测到设备逾期后通知借用人。",
        "params": {
            "card_title": "设备借用逾期",
            "status": "归还逾期，请及时归还，如需继续使用，请到平台点击延期",
            "card_color": "red",
            "status_color": "red",
        },
    },
    "change_borrower_submit_old": {
        "label": "变更申请-原借用人",
        "description": "提交借用人变更申请后通知原借用人。",
        "params": {
            "card_title": "借用人变更申请",
            "status": "待管理员确认，借用人将变为{new_borrower}",
            "card_color": "blue",
            "status_color": "blue",
        },
    },
    "change_borrower_submit_new": {
        "label": "变更申请-新借用人",
        "description": "提交借用人变更申请后通知新借用人。",
        "params": {
            "card_title": "借用人变更申请",
            "status": "待管理员确认，{old_borrower}将变为{new_borrower}",
            "card_color": "blue",
            "status_color": "blue",
        },
    },
    "change_borrower_approve_old": {
        "label": "变更成功-原借用人",
        "description": "管理员确认借用人变更后通知原借用人。",
        "params": {
            "card_title": "借用人变更成功",
            "status": "借用人已变为{new_borrower}",
            "card_color": "green",
            "status_color": "green",
        },
    },
    "change_borrower_approve_new": {
        "label": "变更成功-新借用人",
        "description": "管理员确认借用人变更后通知新借用人。",
        "params": {
            "card_title": "借用人变更成功",
            "status": "借用人已变为{new_borrower}",
            "card_color": "green",
            "status_color": "green",
        },
    },
    "change_borrower_cancel_old": {
        "label": "变更失败-原借用人",
        "description": "管理员取消借用人变更后通知原借用人。",
        "params": {
            "card_title": "借用人变更失败",
            "status": "借用人变更失败",
            "card_color": "red",
            "status_color": "red",
        },
    },
    "change_borrower_cancel_new": {
        "label": "变更失败-新借用人",
        "description": "管理员取消借用人变更后通知新借用人。",
        "params": {
            "card_title": "借用人变更失败",
            "status": "借用人变更失败",
            "card_color": "red",
            "status_color": "red",
        },
    },
}
WEBHOOK_NOTIFICATION_TRIGGER_DEFAULTS: Dict[str, Dict[str, str]] = {
    "borrow_submit": {
        "card_title": "待借通知",
        "body_template": "借用人名字: {borrower}\n借用的设备型号: {device_model}\n归还时间: {return_time}\n待借设备",
    },
    "borrow_approve": {
        "card_title": "借用通知",
        "body_template": (
            "借用人名字: {borrower}\n"
            "借用的设备型号: {device_model}\n"
            "借用时间: {borrow_time}\n"
            "预计归还时间: {return_time}\n"
            "借用成功"
        ),
    },
    "borrow_cancel": {
        "card_title": "借用失败通知",
        "body_template": "借用人名字: {borrower}\n借用的设备型号: {device_model}\n归还时间: {return_time}\n借用失败",
    },
    "extend": {
        "card_title": "延期通知",
        "body_template": (
            "借用人名字: {borrower}\n"
            "借用的设备型号: {device_model}\n"
            "预计归还时间(旧): {old_return_time}\n"
            "变更为 预计归还时间(新): {return_time}"
        ),
    },
    "return": {
        "card_title": "归还通知",
        "body_template": "借用人名字: {borrower}\n借用的设备型号: {device_model}\n归还时间: {return_time}\n归还成功",
    },
    "overdue": {
        "card_title": "逾期通知",
        "body_template": "借用人名字: {borrower}\n借用的设备型号: {device_model}\n预计归还时间: {return_time}\n借用时间已逾期",
    },
    "change_borrower_submit_old": {
        "card_title": "借用人变更通知",
        "body_template": (
            "设备: {device_model}\n"
            "变更前借用人: {old_borrower}\n"
            "变更前归还时间: {old_return_time}\n"
            "变更后借用人: {new_borrower}\n"
            "变更后预期归还时间: {return_time}\n"
            "变更时间: {change_time}"
        ),
    },
    "change_borrower_submit_new": {
        "card_title": "借用人变更通知",
        "body_template": (
            "设备: {device_model}\n"
            "变更前借用人: {old_borrower}\n"
            "变更前归还时间: {old_return_time}\n"
            "变更后借用人: {new_borrower}\n"
            "变更后预期归还时间: {return_time}\n"
            "变更时间: {change_time}"
        ),
    },
    "change_borrower_approve_old": {
        "card_title": "借用人变更成功通知",
        "body_template": "设备: {device_model}\n新的借用人名字: {new_borrower}\n新的归还时间: {return_time}",
    },
    "change_borrower_approve_new": {
        "card_title": "借用人变更成功通知",
        "body_template": "设备: {device_model}\n新的借用人名字: {new_borrower}\n新的归还时间: {return_time}",
    },
    "change_borrower_cancel_old": {
        "card_title": "借用人变更失败通知",
        "body_template": (
            "设备: {device_model}\n"
            "变更前借用人: {old_borrower}\n"
            "变更后借用人: {new_borrower}\n"
            "预计归还时间: {return_time}\n"
            "借用人变更失败"
        ),
    },
    "change_borrower_cancel_new": {
        "card_title": "借用人变更失败通知",
        "body_template": (
            "设备: {device_model}\n"
            "变更前借用人: {old_borrower}\n"
            "变更后借用人: {new_borrower}\n"
            "预计归还时间: {return_time}\n"
            "借用人变更失败"
        ),
    },
}

app = FastAPI()
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

frontend_root = REPO_DIR / "frontend"
frontend_dist = frontend_root / "dist"
frontend_dir = frontend_dist if frontend_dist.exists() else frontend_root
app.mount("/static", StaticFiles(directory=str(frontend_dir)), name="static")


@app.get("/")
async def root():
    return {"message": "Device Loan Assistant API"}


@app.get("/admin")
async def admin_page():
    return FileResponse(str(frontend_dir / "admin.html"))


@app.get("/borrow")
async def borrow_page():
    return FileResponse(str(frontend_dir / "borrow.html"))


def _ensure_required(value: Optional[str], label: str):
    if value is None or not str(value).strip():
        raise HTTPException(status_code=400, detail=f"{label}不能为空")


def _parse_datetime(value: str) -> datetime:
    try:
        cleaned = value.replace("Z", "+00:00") if value.endswith("Z") else value
        parsed = datetime.fromisoformat(cleaned)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"时间格式错误: {exc}")
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed


LOCAL_TZ = timezone(timedelta(hours=8))


def _format_notify_time(value: Optional[Union[str, datetime]]) -> str:
    if value is None:
        return "-"
    if isinstance(value, datetime):
        parsed = value
    else:
        parsed = _parse_datetime(value)
    return parsed.astimezone(LOCAL_TZ).strftime("%Y-%m-%d %H:%M:%S")


def _format_loan_status(value: Optional[str]) -> str:
    mapping = {"available": "可借", "pending": "待借", "borrowed": "已借"}
    if not value:
        return "-"
    return mapping.get(value, value)


def _extract_performance(notes: Optional[str]) -> str:
    if not notes:
        return "-"
    cleaned = " ".join(notes.split()).strip()
    if not cleaned:
        return "-"
    match = re.search(r"性能[:：=\s]*([^\n,，;；。]+)", cleaned)
    scope = match.group(1) if match else cleaned
    keywords = ["强劲", "较高", "一般", "较低", "高", "中", "低", "强", "弱"]
    for keyword in keywords:
        if keyword in scope:
            return keyword
    if match:
        value = match.group(1).strip()
        return value or "-"
    return "-"


def _row_to_dict(row) -> Dict[str, Any]:
    return dict(row) if row else {}


def _get_setting(conn, key: str) -> str:
    row = conn.execute("SELECT value FROM settings WHERE key = ?", (key,)).fetchone()
    if row:
        return row["value"]
    return ""


def _set_setting(conn, key: str, value: str) -> None:
    conn.execute(
        "INSERT INTO settings (key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = excluded.value",
        (key, value),
    )


def _delete_setting(conn, key: str) -> None:
    conn.execute("DELETE FROM settings WHERE key = ?", (key,))


def _clean_text(value: Any) -> str:
    return str(value or "").strip()


def _default_notification_params(trigger: str) -> Dict[str, str]:
    config = NOTIFICATION_TRIGGER_DEFAULTS.get(trigger)
    if not config:
        raise HTTPException(status_code=400, detail="未知通知触发点")
    return dict(config["params"])


def _normalize_notification_params(values: Dict[str, Any]) -> Dict[str, str]:
    cleaned: Dict[str, str] = {}
    for field in NOTIFICATION_PARAM_FIELDS:
        value = _clean_text(values.get(field))
        if not value:
            raise HTTPException(status_code=400, detail="通知参数不能为空")
        if field in NOTIFICATION_COLOR_FIELDS and value not in NOTIFICATION_COLOR_OPTIONS:
            raise HTTPException(status_code=400, detail="通知颜色参数无效")
        cleaned[field] = value
    return cleaned


def _load_notification_param_overrides(conn) -> Dict[str, Dict[str, str]]:
    value = _get_setting(conn, NOTIFICATION_SETTINGS_KEY)
    if not value:
        return {}
    try:
        raw = json.loads(value)
    except json.JSONDecodeError:
        return {}
    if not isinstance(raw, dict):
        return {}
    overrides: Dict[str, Dict[str, str]] = {}
    for trigger, params in raw.items():
        if trigger not in NOTIFICATION_TRIGGER_DEFAULTS or not isinstance(params, dict):
            continue
        try:
            overrides[trigger] = _normalize_notification_params(params)
        except HTTPException:
            continue
    return overrides


def _notification_params_for_trigger(
    trigger: str,
    overrides: Dict[str, Dict[str, str]],
    template_values: Optional[Dict[str, Any]] = None,
) -> Dict[str, str]:
    params = _default_notification_params(trigger)
    params.update(overrides.get(trigger, {}))
    if template_values:
        params["card_title"] = _render_notification_template(params["card_title"], template_values)
        params["status"] = _render_notification_template(params["status"], template_values)
    return params


def _render_notification_template(value: str, values: Dict[str, Any]) -> str:
    def replace(match: re.Match[str]) -> str:
        return _clean_text(values.get(match.group(1))) or "-"

    return re.sub(r"\{([A-Za-z0-9_]+)\}", replace, value)


def _notification_settings_items(overrides: Dict[str, Dict[str, str]]) -> List[Dict[str, Any]]:
    items = []
    for trigger, config in NOTIFICATION_TRIGGER_DEFAULTS.items():
        params = _notification_params_for_trigger(trigger, overrides)
        items.append(
            {
                "key": trigger,
                "label": config["label"],
                "description": config["description"],
                "defaults": _default_notification_params(trigger),
                "params": params,
                "customized": params != _default_notification_params(trigger),
            }
        )
    return items


def _save_notification_settings(conn, settings: Dict[str, Any]) -> Dict[str, Dict[str, str]]:
    overrides: Dict[str, Dict[str, str]] = {}
    for trigger, params in settings.items():
        if trigger not in NOTIFICATION_TRIGGER_DEFAULTS:
            raise HTTPException(status_code=400, detail="未知通知触发点")
        raw_params = params.model_dump() if hasattr(params, "model_dump") else params.dict()
        normalized = _normalize_notification_params(raw_params)
        if normalized != _default_notification_params(trigger):
            overrides[trigger] = normalized
    _set_setting(conn, NOTIFICATION_SETTINGS_KEY, json.dumps(overrides, ensure_ascii=False, sort_keys=True))
    return overrides


def _default_webhook_notification_params(trigger: str) -> Dict[str, str]:
    defaults = WEBHOOK_NOTIFICATION_TRIGGER_DEFAULTS.get(trigger)
    if not defaults:
        raise HTTPException(status_code=400, detail="未知通知触发点")
    return {
        "card_title": defaults["card_title"],
        "body_template": defaults["body_template"],
        "card_color": _default_notification_params(trigger)["card_color"],
    }


def _normalize_webhook_notification_params(values: Dict[str, Any]) -> Dict[str, str]:
    cleaned: Dict[str, str] = {}
    for field in WEBHOOK_NOTIFICATION_PARAM_FIELDS:
        value = _clean_text(values.get(field))
        if not value:
            raise HTTPException(status_code=400, detail="Webhook 通知参数不能为空")
        if field == "card_color" and value not in NOTIFICATION_COLOR_OPTIONS:
            raise HTTPException(status_code=400, detail="Webhook 通知颜色参数无效")
        cleaned[field] = value
    return cleaned


def _load_webhook_notification_param_overrides(conn) -> Dict[str, Dict[str, str]]:
    value = _get_setting(conn, WEBHOOK_NOTIFICATION_SETTINGS_KEY)
    if not value:
        return {}
    try:
        raw = json.loads(value)
    except json.JSONDecodeError:
        return {}
    if not isinstance(raw, dict):
        return {}
    overrides: Dict[str, Dict[str, str]] = {}
    for trigger, params in raw.items():
        if trigger not in NOTIFICATION_TRIGGER_DEFAULTS or not isinstance(params, dict):
            continue
        try:
            overrides[trigger] = _normalize_webhook_notification_params(params)
        except HTTPException:
            continue
    return overrides


def _webhook_notification_params_for_trigger(
    trigger: str, overrides: Dict[str, Dict[str, str]]
) -> Dict[str, str]:
    params = _default_webhook_notification_params(trigger)
    params.update(overrides.get(trigger, {}))
    return params


def _webhook_notification_settings_items(overrides: Dict[str, Dict[str, str]]) -> List[Dict[str, Any]]:
    items = []
    for trigger, config in NOTIFICATION_TRIGGER_DEFAULTS.items():
        params = _webhook_notification_params_for_trigger(trigger, overrides)
        defaults = _default_webhook_notification_params(trigger)
        items.append(
            {
                "key": trigger,
                "label": config["label"],
                "description": config["description"],
                "defaults": defaults,
                "params": params,
                "customized": params != defaults,
            }
        )
    return items


def _save_webhook_notification_settings(conn, settings: Dict[str, Any]) -> Dict[str, Dict[str, str]]:
    overrides: Dict[str, Dict[str, str]] = {}
    for trigger, params in settings.items():
        if trigger not in NOTIFICATION_TRIGGER_DEFAULTS:
            raise HTTPException(status_code=400, detail="未知通知触发点")
        raw_params = params.model_dump() if hasattr(params, "model_dump") else params.dict()
        normalized = _normalize_webhook_notification_params(raw_params)
        if normalized != _default_webhook_notification_params(trigger):
            overrides[trigger] = normalized
    _set_setting(conn, WEBHOOK_NOTIFICATION_SETTINGS_KEY, json.dumps(overrides, ensure_ascii=False, sort_keys=True))
    return overrides


def _borrow_admin_url(conn) -> str:
    return _get_setting(conn, BORROW_ADMIN_URL_SETTING_KEY).strip() or DEFAULT_BORROW_ADMIN_URL


def _webhook_card_color_for_trigger(trigger: Optional[str], conn) -> str:
    if not trigger or trigger not in NOTIFICATION_TRIGGER_DEFAULTS:
        return "blue"
    overrides = _load_webhook_notification_param_overrides(conn)
    return _webhook_notification_params_for_trigger(trigger, overrides)["card_color"]


def _build_configured_webhook_message(
    conn,
    *,
    trigger: Optional[str],
    fallback_title: str,
    fallback_body: str,
    template_values: Optional[Dict[str, Any]] = None,
) -> Dict[str, str]:
    admin_url = _borrow_admin_url(conn)
    if not trigger or trigger not in NOTIFICATION_TRIGGER_DEFAULTS:
        return {"title": fallback_title, "body": fallback_body, "card_color": "blue", "admin_url": admin_url}
    overrides = _load_webhook_notification_param_overrides(conn)
    params = _webhook_notification_params_for_trigger(trigger, overrides)
    values = template_values
    title = params["card_title"]
    body = params["body_template"]
    if values is not None:
        title = _render_notification_template(title, values)
        body = _render_notification_template(body, values)
    return {
        "title": title,
        "body": body,
        "card_color": params["card_color"],
        "admin_url": admin_url,
    }


def _legacy_borrower_profile(name: Optional[str]) -> Dict[str, Optional[str]]:
    return {
        "name": _clean_text(name),
        "user_id": None,
        "open_id": None,
        "avatar_url": None,
        "job_title": None,
    }


def _borrower_profile_from_user(user: Dict[str, Any]) -> Dict[str, Optional[str]]:
    name = _clean_text(user.get("name")) or "已登录用户"
    return {
        "name": name,
        "user_id": _clean_text(user.get("user_id")) or None,
        "open_id": _clean_text(user.get("open_id")) or None,
        "avatar_url": _clean_text(user.get("avatar_url")) or None,
        "job_title": _clean_text(user.get("job_title")) or None,
    }


def _borrower_profile_from_row(row: Union[Dict[str, Any], Any], prefix: str = "borrower") -> Dict[str, Optional[str]]:
    if row is None:
        return _legacy_borrower_profile(None)
    getter = row.get if isinstance(row, dict) else lambda key, default=None: row[key] if key in row.keys() else default
    return {
        "name": _clean_text(getter(f"{prefix}_name", None) if prefix == "borrower" else getter(prefix, None)),
        "user_id": _clean_text(getter(f"{prefix}_user_id", None)) or None,
        "open_id": _clean_text(getter(f"{prefix}_open_id", None)) or None,
        "avatar_url": _clean_text(getter(f"{prefix}_avatar_url", None)) or None,
        "job_title": _clean_text(getter(f"{prefix}_job_title", None)) or None,
    }


def _borrower_values(profile: Dict[str, Optional[str]]) -> tuple:
    return (
        profile.get("name"),
        profile.get("user_id"),
        profile.get("open_id"),
        profile.get("avatar_url"),
        profile.get("job_title"),
    )


def _get_portal_token(request: Request) -> str:
    auth_header = request.headers.get("authorization", "")
    if auth_header.lower().startswith("bearer "):
        return auth_header[7:].strip()
    return request.headers.get("x-portal-jwt", "").strip()


async def _require_portal_borrower_profile(request: Request) -> Dict[str, Optional[str]]:
    token = _get_portal_token(request)
    if not token:
        raise HTTPException(status_code=401, detail="缺少门户登录凭证", headers={"X-Auth-Error": "SSO_JWT_MISSING"})

    audience = request.headers.get("x-portal-audience") or request.headers.get("origin") or str(request.base_url).rstrip("/")
    client: httpx.AsyncClient = app.state.http_client
    try:
        response = await client.post(
            PORTAL_JWT_VERIFY_URL,
            json={"token": token, "audience": audience},
        )
    except httpx.HTTPError:
        raise HTTPException(status_code=503, detail="门户登录校验暂不可用")

    try:
        payload = response.json()
    except ValueError:
        payload = {}

    if response.status_code >= 400 or not payload.get("success"):
        code = payload.get("code") or "SSO_JWT_INVALID"
        detail = "登录凭证已失效，正在重新登录/请重新登录" if code == "SSO_JWT_INVALID" else payload.get("message", "门户登录校验失败")
        raise HTTPException(status_code=401, detail=f"{code}: {detail}", headers={"X-Auth-Error": str(code)})

    data = payload.get("data") or {}
    user = data.get("user") or {}
    if not isinstance(user, dict):
        raise HTTPException(status_code=401, detail="SSO_JWT_INVALID: 登录凭证已失效，正在重新登录/请重新登录")
    return _borrower_profile_from_user(user)


async def _resolve_borrower_profile(request: Request, fallback_name: Optional[str]) -> Dict[str, Optional[str]]:
    if _get_portal_token(request):
        return await _require_portal_borrower_profile(request)
    _ensure_required(fallback_name, "借用人名字")
    return _legacy_borrower_profile(fallback_name)


def _portal_notification_auth_from_request(request: Request) -> Dict[str, Optional[str]]:
    token = _get_portal_token(request)
    if token:
        return {"portal_jwt": token, "cookie": None}
    cookie = request.headers.get("cookie", "").strip()
    return {"portal_jwt": None, "cookie": cookie or None}


def _portal_notification_auth_from_job() -> Dict[str, Optional[str]]:
    service_id = os.environ.get(PORTAL_NOTIFICATION_SERVICE_ID_ENV, DEFAULT_PORTAL_NOTIFICATION_SERVICE_ID).strip()
    service_token = os.environ.get(PORTAL_NOTIFICATION_SERVICE_TOKEN_ENV, "").strip()
    return {
        "service_id": service_id or DEFAULT_PORTAL_NOTIFICATION_SERVICE_ID,
        "service_token": service_token or None,
        "source": "service" if service_token else "none",
    }


def _sanitize_error_message(value: Any) -> str:
    message = str(value or "").strip()
    message = re.sub(r"Bearer\s+[^\s,;]+", "Bearer <redacted>", message, flags=re.IGNORECASE)
    message = re.sub(r"(Authorization\s*[:=]\s*)[^\s,;]+", r"\1<redacted>", message, flags=re.IGNORECASE)
    message = re.sub(r"(Cookie\s*[:=]\s*)[^,;]+", r"\1<redacted>", message, flags=re.IGNORECASE)
    message = re.sub(r"(portal_jwt\s*[:=]\s*)[^\s,;]+", r"\1<redacted>", message, flags=re.IGNORECASE)
    return message[:300] if message else "门户通知发送失败"


def _portal_error_parts(exc: PortalNotificationError) -> tuple:
    return exc.status_code, exc.code, _sanitize_error_message(exc.message)


def _log_portal_notification_error(context: str, exc: PortalNotificationError) -> None:
    status_code, code, message = _portal_error_parts(exc)
    logger.warning(
        "门户通知发送失败: context=%s status=%s code=%s message=%s",
        context,
        status_code,
        code,
        message,
    )


def _build_portal_card_payload(
    *,
    borrower: Optional[str],
    device_name: Optional[str],
    request_date: Optional[Union[str, datetime]],
    return_date: Optional[Union[str, datetime]],
    status: str,
    card_color: str,
    status_color: str,
    card_title: str,
) -> Dict[str, Any]:
    return {
        "borrower": borrower or "-",
        "device_name": device_name or "-",
        "request_date": _format_notify_time(request_date),
        "return_date": _format_notify_time(return_date),
        "status": status,
        "card_color": card_color,
        "status_color": status_color,
        "card_title": card_title,
    }


def _build_configured_portal_card_payload(
    *,
    trigger: str,
    borrower: Optional[str],
    device_name: Optional[str],
    request_date: Optional[Union[str, datetime]],
    return_date: Optional[Union[str, datetime]],
    template_values: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    with db_session() as conn:
        overrides = _load_notification_param_overrides(conn)
    params = _notification_params_for_trigger(trigger, overrides, template_values)
    return _build_portal_card_payload(
        borrower=borrower,
        device_name=device_name,
        request_date=request_date,
        return_date=return_date,
        status=params["status"],
        card_color=params["card_color"],
        status_color=params["status_color"],
        card_title=params["card_title"],
    )


async def _queue_portal_notification(
    context: str,
    recipient_user_id: Optional[str],
    payload: Dict[str, Any],
    portal_jwt: Optional[str],
    cookie: Optional[str],
) -> None:
    client: httpx.AsyncClient = app.state.http_client
    try:
        await send_portal_notification(
            client,
            recipient_user_id=recipient_user_id or "",
            payload=payload,
            portal_jwt=portal_jwt,
            cookie=cookie,
        )
    except PortalNotificationError as exc:
        _log_portal_notification_error(context, exc)
    except Exception:
        logger.warning("门户通知发送失败: context=%s status=%s code=%s message=%s", context, None, "UNEXPECTED_ERROR", "未知错误")


def _add_portal_notification_task(
    background_tasks: BackgroundTasks,
    *,
    context: str,
    recipient_user_id: Optional[str],
    payload: Dict[str, Any],
    auth: Dict[str, Optional[str]],
) -> None:
    background_tasks.add_task(
        _queue_portal_notification,
        context,
        recipient_user_id,
        payload,
        auth.get("portal_jwt"),
        auth.get("cookie"),
    )


def _fetch_current_borrow_context(conn, device_id: int) -> Dict[str, Any]:
    record = conn.execute(
        """
        SELECT br.id AS borrow_record_id, br.request_id, br.borrowed_at, br.expected_return_at,
               br.borrower_name, br.borrower_user_id, rq.requested_at
        FROM borrow_records br
        LEFT JOIN borrow_requests rq ON rq.id = br.request_id
        WHERE br.device_id = ? AND br.returned_at IS NULL
        ORDER BY br.borrowed_at DESC
        LIMIT 1
        """,
        (device_id,),
    ).fetchone()
    if record:
        return dict(record)
    request = conn.execute(
        """
        SELECT NULL AS borrow_record_id, br.id AS request_id, NULL AS borrowed_at, br.expected_return_at,
               br.borrower_name, br.borrower_user_id, br.requested_at
        FROM borrow_requests br
        WHERE br.device_id = ? AND br.request_type = 'borrow' AND br.status = 'pending'
        ORDER BY br.requested_at DESC
        LIMIT 1
        """,
        (device_id,),
    ).fetchone()
    return dict(request) if request else {}


def _borrow_context_request_date(context: Dict[str, Any], fallback: Optional[Union[str, datetime]] = None) -> Optional[Union[str, datetime]]:
    return context.get("requested_at") or context.get("borrowed_at") or fallback


def _get_int_setting(conn, key: str) -> Optional[int]:
    value = _get_setting(conn, key)
    if value and value.isdigit():
        return int(value)
    return None


def _insert_borrow_record(
    conn,
    device_id: int,
    device_model: str,
    borrower_profile: Dict[str, Optional[str]],
    borrowed_at: str,
    expected_return_at: Optional[str],
    request_id: Optional[int],
) -> None:
    now = now_iso()
    conn.execute(
        """
        INSERT INTO borrow_records (
            device_id, device_model, borrower_name, borrower_user_id, borrower_open_id,
            borrower_avatar_url, borrower_job_title, borrowed_at, expected_return_at,
            returned_at, status, request_id, created_at, updated_at
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, 'borrowed', ?, ?, ?)
        """,
        (
            device_id,
            device_model,
            *_borrower_values(borrower_profile),
            borrowed_at,
            expected_return_at,
            request_id,
            now,
            now,
        ),
    )


def _update_borrow_record_expected(conn, device_id: int, expected_return_at: str) -> None:
    now = now_iso()
    conn.execute(
        """
        UPDATE borrow_records
        SET expected_return_at = ?, updated_at = ?
        WHERE id = (
            SELECT id FROM borrow_records
            WHERE device_id = ? AND returned_at IS NULL
            ORDER BY borrowed_at DESC
            LIMIT 1
        )
        """,
        (expected_return_at, now, device_id),
    )


def _close_borrow_record(conn, device_id: int, returned_at: str) -> None:
    now = now_iso()
    conn.execute(
        """
        UPDATE borrow_records
        SET returned_at = ?, status = 'returned', updated_at = ?
        WHERE id = (
            SELECT id FROM borrow_records
            WHERE device_id = ? AND returned_at IS NULL
            ORDER BY borrowed_at DESC
            LIMIT 1
        )
        """,
        (returned_at, now, device_id),
    )


def _assign_llm_model(conn, model_id: int, role: str) -> None:
    role_value = role.strip().lower()
    if role_value not in ("fast", "accurate"):
        raise HTTPException(status_code=400, detail="角色必须为 fast 或 accurate")
    row = conn.execute("SELECT id FROM llm_models WHERE id = ?", (model_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="模型不存在")
    fast_id = _get_int_setting(conn, "llm_fast_model_id")
    accurate_id = _get_int_setting(conn, "llm_accurate_model_id")
    if role_value == "fast":
        if accurate_id == model_id:
            raise HTTPException(status_code=400, detail="更快模型不能与更准模型相同")
        _set_setting(conn, "llm_fast_model_id", str(model_id))
    else:
        if fast_id == model_id:
            raise HTTPException(status_code=400, detail="更准模型不能与更快模型相同")
        _set_setting(conn, "llm_accurate_model_id", str(model_id))


def _clear_llm_assignment(conn, role: str) -> None:
    role_value = role.strip().lower()
    if role_value not in ("fast", "accurate"):
        raise HTTPException(status_code=400, detail="角色必须为 fast 或 accurate")
    key = "llm_fast_model_id" if role_value == "fast" else "llm_accurate_model_id"
    _delete_setting(conn, key)


def _validate_version_for_system(conn, system_id: int, version_id: int) -> None:
    row = conn.execute(
        "SELECT id FROM system_versions WHERE id = ? AND system_id = ?",
        (version_id, system_id),
    ).fetchone()
    if not row:
        raise HTTPException(status_code=400, detail="系统版本不匹配")


def _fetch_devices(conn, query: Optional[str] = None) -> List[Dict[str, Any]]:
    sql = (
        "SELECT d.*, v.name AS vendor_name, s.name AS system_name, sv.version AS system_version "
        "FROM devices d "
        "JOIN vendors v ON d.vendor_id = v.id "
        "JOIN systems s ON d.system_id = s.id "
        "JOIN system_versions sv ON d.system_version_id = sv.id "
    )
    params: List[Any] = []
    if query:
        like = f"%{query}%"
        sql += (
            "WHERE CAST(d.id AS TEXT) LIKE ? OR d.model LIKE ? OR d.status LIKE ? OR d.type LIKE ? OR v.name LIKE ? "
            "OR s.name LIKE ? OR sv.version LIKE ? OR d.resolution LIKE ? OR d.arch LIKE ? "
            "OR d.cpu LIKE ? OR d.notes LIKE ? OR d.borrower_name LIKE ? "
        )
        params = [
            like,
            like,
            like,
            like,
            like,
            like,
            like,
            like,
            like,
            like,
            like,
            like,
        ]
    sql += "ORDER BY d.id DESC"
    rows = conn.execute(sql, params).fetchall()
    return [dict(row) for row in rows]


def _fetch_borrow_requests(
    conn,
    query: Optional[str] = None,
    status: Optional[str] = None,
    device_id: Optional[int] = None,
    request_type: Optional[str] = None,
) -> List[Dict[str, Any]]:
    sql = (
        "SELECT "
        "br.id, br.device_id, br.device_model, br.borrower_name, br.borrower_user_id, "
        "br.borrower_open_id, br.borrower_avatar_url, br.borrower_job_title, br.expected_return_at, "
        "br.request_type, br.status AS request_status, br.requested_at, br.handled_at, "
        "d.status AS device_status, d.type AS device_type, v.name AS vendor_name, "
        "s.name AS system_name, sv.version AS system_version, "
        "d.resolution, d.arch, d.cpu, d.boot_password, d.notes "
        "FROM borrow_requests br "
        "LEFT JOIN devices d ON br.device_id = d.id "
        "LEFT JOIN vendors v ON d.vendor_id = v.id "
        "LEFT JOIN systems s ON d.system_id = s.id "
        "LEFT JOIN system_versions sv ON d.system_version_id = sv.id "
    )
    params: List[Any] = []
    conditions: List[str] = []
    if query:
        like = f"%{query}%"
        conditions.append("(br.borrower_name LIKE ? OR br.device_model LIKE ? OR br.request_type LIKE ?)")
        params.extend([like, like, like])
    if status:
        conditions.append("br.status = ?")
        params.append(status)
    if device_id is not None:
        conditions.append("br.device_id = ?")
        params.append(device_id)
    if request_type:
        conditions.append("br.request_type = ?")
        params.append(request_type)
    if conditions:
        sql += "WHERE " + " AND ".join(conditions) + " "
    sql += "ORDER BY br.requested_at DESC"
    rows = conn.execute(sql, params).fetchall()
    return [dict(row) for row in rows]


def _fetch_borrow_records(conn, query: Optional[str] = None) -> List[Dict[str, Any]]:
    sql = (
        "SELECT id, device_id, device_model, borrower_name, borrower_user_id, borrower_open_id, "
        "borrower_avatar_url, borrower_job_title, borrowed_at, expected_return_at, "
        "returned_at, status, request_id, "
        "("
        "SELECT manual_sent_at FROM overdue_notifications od "
        "WHERE od.borrow_record_id = borrow_records.id "
        "AND od.expected_return_at = borrow_records.expected_return_at "
        "AND od.manual_sent_at IS NOT NULL "
        "ORDER BY od.manual_sent_at DESC LIMIT 1"
        ") AS overdue_manual_sent_at "
        "FROM borrow_records "
    )
    params: List[Any] = []
    if query:
        like = f"%{query}%"
        sql += "WHERE borrower_name LIKE ? OR device_model LIKE ? OR status LIKE ? "
        params = [like, like, like]
    sql += "ORDER BY borrowed_at DESC"
    rows = conn.execute(sql, params).fetchall()
    items = [dict(row) for row in rows]
    if not items:
        return items
    record_ids = [item["id"] for item in items if item.get("id") is not None]
    request_ids = [item["request_id"] for item in items if item.get("request_id") is not None]
    if not record_ids and not request_ids:
        return items
    conditions = []
    params: List[Any] = []
    if record_ids:
        placeholders = ",".join("?" for _ in record_ids)
        conditions.append(f"record_id IN ({placeholders})")
        params.extend(record_ids)
    if request_ids:
        placeholders = ",".join("?" for _ in request_ids)
        conditions.append(f"request_id IN ({placeholders})")
        params.extend(request_ids)
    sql = (
        "SELECT id, record_id, request_id, borrower_before, borrower_before_user_id, "
        "borrower_before_open_id, borrower_before_avatar_url, borrower_before_job_title, "
        "borrower_after, borrower_after_user_id, borrower_after_open_id, borrower_after_avatar_url, "
        "borrower_after_job_title, expected_before, "
        "expected_after, changed_at FROM borrow_changes "
    )
    if conditions:
        sql += "WHERE " + " OR ".join(conditions) + " "
    sql += "ORDER BY changed_at ASC"
    change_rows = conn.execute(sql, params).fetchall()
    change_by_record: Dict[int, List[Dict[str, Any]]] = {}
    change_by_request: Dict[int, List[Dict[str, Any]]] = {}
    for row in change_rows:
        data = dict(row)
        if row["record_id"] is not None:
            change_by_record.setdefault(row["record_id"], []).append(data)
        if row["request_id"] is not None:
            change_by_request.setdefault(row["request_id"], []).append(data)
    for item in items:
        record_list = change_by_record.get(item["id"], [])
        request_list = change_by_request.get(item.get("request_id"), [])
        if record_list or request_list:
            item["borrower_changes"] = record_list + [c for c in request_list if c not in record_list]
        else:
            item["borrower_changes"] = []
    return items


def _migrate_borrower_table(conn, table: str, profile: Dict[str, Optional[str]]) -> int:
    name = profile["name"]
    now = now_iso()
    cur = conn.execute(
        f"""
        UPDATE {table}
        SET borrower_name = ?,
            borrower_user_id = ?,
            borrower_open_id = ?,
            borrower_avatar_url = ?,
            borrower_job_title = ?,
            updated_at = ?
        WHERE borrower_name IS NOT NULL
          AND TRIM(borrower_name) = ?
          AND (
            borrower_name <> ?
            OR COALESCE(borrower_user_id, '') <> COALESCE(?, '')
            OR COALESCE(borrower_open_id, '') <> COALESCE(?, '')
            OR COALESCE(borrower_avatar_url, '') <> COALESCE(?, '')
            OR COALESCE(borrower_job_title, '') <> COALESCE(?, '')
          )
        """,
        (
            *_borrower_values(profile),
            now,
            name,
            profile.get("name"),
            profile.get("user_id"),
            profile.get("open_id"),
            profile.get("avatar_url"),
            profile.get("job_title"),
        ),
    )
    return cur.rowcount


def _migrate_borrower_change_side(conn, side: str, profile: Dict[str, Optional[str]]) -> int:
    name = profile["name"]
    now = now_iso()
    cur = conn.execute(
        f"""
        UPDATE borrow_changes
        SET {side} = ?,
            {side}_user_id = ?,
            {side}_open_id = ?,
            {side}_avatar_url = ?,
            {side}_job_title = ?,
            updated_at = ?
        WHERE {side} IS NOT NULL
          AND TRIM({side}) = ?
          AND (
            {side} <> ?
            OR COALESCE({side}_user_id, '') <> COALESCE(?, '')
            OR COALESCE({side}_open_id, '') <> COALESCE(?, '')
            OR COALESCE({side}_avatar_url, '') <> COALESCE(?, '')
            OR COALESCE({side}_job_title, '') <> COALESCE(?, '')
          )
        """,
        (
            *_borrower_values(profile),
            now,
            name,
            profile.get("name"),
            profile.get("user_id"),
            profile.get("open_id"),
            profile.get("avatar_url"),
            profile.get("job_title"),
        ),
    )
    return cur.rowcount


def _filter_ai_devices(devices: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    return devices


def _upsert_pending_overdue_notifications(conn, now: datetime) -> None:
    rows = conn.execute(
        """
        SELECT d.id AS device_id, d.model AS device_model, d.borrower_name, d.borrower_user_id,
               d.expected_return_at, br.id AS borrow_record_id, br.borrowed_at, rq.requested_at
        FROM devices d
        JOIN borrow_records br ON br.id = (
            SELECT id FROM borrow_records
            WHERE device_id = d.id AND returned_at IS NULL
            ORDER BY borrowed_at DESC
            LIMIT 1
        )
        LEFT JOIN borrow_requests rq ON rq.id = br.request_id
        WHERE d.loan_status = 'borrowed'
          AND d.expected_return_at IS NOT NULL
          AND d.overdue_notified = 0
          AND d.status <> ?
        """,
        (RESIDENT_DEVICE_STATUS,),
    ).fetchall()
    for row in rows:
        try:
            expected = _parse_datetime(row["expected_return_at"])
        except HTTPException:
            continue
        if expected >= now:
            continue
        borrower_user_id = _clean_text(row["borrower_user_id"])
        borrower_name = _clean_text(row["borrower_name"])
        current = now_iso()
        conn.execute(
            """
            INSERT OR IGNORE INTO overdue_notifications (
                device_id, borrow_record_id, borrower_user_id, borrower_name, device_model,
                requested_at, expected_return_at, status, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?)
            """,
            (
                row["device_id"],
                row["borrow_record_id"],
                borrower_user_id,
                borrower_name,
                row["device_model"],
                row["requested_at"] or row["borrowed_at"],
                row["expected_return_at"],
                current,
                current,
            ),
        )


def _is_overdue_snapshot_current(conn, notification: Dict[str, Any], now: datetime, *, require_recipient: bool = True) -> bool:
    recipient = _clean_text(notification.get("borrower_user_id"))
    borrower_name = _clean_text(notification.get("borrower_name"))
    if not borrower_name:
        return False
    if require_recipient and not recipient:
        return False
    try:
        expected = _parse_datetime(notification["expected_return_at"])
    except HTTPException:
        return False
    if expected >= now:
        return False
    current = conn.execute(
        """
        SELECT d.id
        FROM devices d
        JOIN borrow_records br ON br.id = ? AND br.device_id = d.id AND br.returned_at IS NULL
        WHERE d.id = ?
          AND d.loan_status = 'borrowed'
          AND d.status <> ?
          AND d.expected_return_at = ?
          AND d.borrower_name IS NOT NULL
          AND TRIM(d.borrower_name) <> ''
          AND d.borrower_name = ?
        """,
        (
            notification["borrow_record_id"],
            notification["device_id"],
            RESIDENT_DEVICE_STATUS,
            notification["expected_return_at"],
            borrower_name,
        ),
    ).fetchone()
    if current is None:
        return False
    if require_recipient:
        device_recipient = conn.execute(
            "SELECT COALESCE(borrower_user_id, '') AS borrower_user_id FROM devices WHERE id = ?",
            (notification["device_id"],),
        ).fetchone()
        return device_recipient is not None and device_recipient["borrower_user_id"] == recipient
    return True


def _mark_overdue_notification_skipped(conn, notification_id: int, code: str, message: str) -> None:
    conn.execute(
        """
        UPDATE overdue_notifications
        SET status = 'skipped', last_error_code = ?, last_error_message = ?,
            last_error_status = NULL, updated_at = ?
        WHERE id = ? AND status = 'pending'
        """,
        (code, _sanitize_error_message(message), now_iso(), notification_id),
    )


def _record_overdue_notification_error(conn, notification_id: int, exc: PortalNotificationError) -> None:
    status_code, code, message = _portal_error_parts(exc)
    next_status = "failed" if status_code == 400 else "pending"
    conn.execute(
        """
        UPDATE overdue_notifications
        SET status = ?, last_error_code = ?, last_error_message = ?,
            last_error_status = ?, updated_at = ?
        WHERE id = ? AND status = 'pending'
        """,
        (next_status, code, message, status_code, now_iso(), notification_id),
    )


def _record_overdue_auth_missing(conn) -> None:
    conn.execute(
        """
        UPDATE overdue_notifications
        SET last_error_code = 'NOTIFICATION_AUTH_REQUIRED',
            last_error_message = '缺少门户通知定时任务鉴权凭证',
            last_error_status = 401,
            updated_at = ?
        WHERE status = 'pending'
        """,
        (now_iso(),),
    )


def _pending_overdue_webhook_notifications(conn, now: datetime) -> List[Dict[str, Any]]:
    rows = conn.execute(
        """
        SELECT id, device_id, borrow_record_id, borrower_user_id, borrower_name,
               device_model, requested_at, expected_return_at
        FROM overdue_notifications
        WHERE status = 'pending' AND webhook_sent_at IS NULL
        ORDER BY created_at ASC
        """
    ).fetchall()
    pending = []
    for row in rows:
        item = dict(row)
        if _is_overdue_snapshot_current(conn, item, now, require_recipient=False):
            pending.append(item)
            continue
        _mark_overdue_notification_skipped(
            conn,
            item["id"],
            "OVERDUE_NOTIFICATION_STALE",
            "逾期通知快照已失效",
        )
    return pending


def _mark_overdue_webhook_sent(conn, notification: Dict[str, Any], sent_at: str) -> None:
    conn.execute(
        """
        UPDATE overdue_notifications
        SET webhook_sent_at = ?, webhook_last_error_message = NULL, updated_at = ?
        WHERE id = ? AND webhook_sent_at IS NULL
        """,
        (sent_at, sent_at, notification["id"]),
    )
    conn.execute(
        "UPDATE devices SET overdue_notified = 1, updated_at = ? WHERE id = ?",
        (sent_at, notification["device_id"]),
    )


def _record_overdue_webhook_error(conn, notification_id: int, message: str) -> None:
    conn.execute(
        """
        UPDATE overdue_notifications
        SET webhook_last_error_message = ?, updated_at = ?
        WHERE id = ? AND webhook_sent_at IS NULL
        """,
        (_sanitize_error_message(message), now_iso(), notification_id),
    )


def _mark_overdue_notification_sent(conn, notification: Dict[str, Any], sent_at: str) -> bool:
    now = _parse_datetime(sent_at)
    if not _is_overdue_snapshot_current(conn, notification, now):
        _mark_overdue_notification_skipped(
            conn,
            notification["id"],
            "OVERDUE_NOTIFICATION_STALE",
            "逾期通知快照已失效",
        )
        return False
    cur = conn.execute(
        """
        UPDATE overdue_notifications
        SET status = 'sent', sent_at = ?, last_error_code = NULL, last_error_message = NULL,
            last_error_status = NULL, updated_at = ?
        WHERE id = ? AND status = 'pending'
        """,
        (sent_at, sent_at, notification["id"]),
    )
    if cur.rowcount == 0:
        return False
    conn.execute(
        "UPDATE devices SET overdue_notified = 1, updated_at = ? WHERE id = ?",
        (sent_at, notification["device_id"]),
    )
    return True


def _pending_overdue_notifications(conn, now: datetime) -> List[Dict[str, Any]]:
    rows = conn.execute(
        """
        SELECT id, device_id, borrow_record_id, borrower_user_id, borrower_name,
               device_model, requested_at, expected_return_at
        FROM overdue_notifications
        WHERE status = 'pending'
        ORDER BY created_at ASC
        """
    ).fetchall()
    pending = []
    for row in rows:
        item = dict(row)
        if _is_overdue_snapshot_current(conn, item, now):
            pending.append(item)
            continue
        _mark_overdue_notification_skipped(
            conn,
            item["id"],
            "OVERDUE_NOTIFICATION_STALE",
            "逾期通知快照已失效",
        )
    return pending


def _fetch_borrow_record_overdue_context(conn, record_id: int) -> Dict[str, Any]:
    row = conn.execute(
        """
        SELECT br.id AS borrow_record_id, br.device_id, br.device_model,
               br.borrower_name, br.borrower_user_id, br.borrowed_at,
               br.expected_return_at, br.returned_at, br.status AS record_status,
               d.loan_status AS device_loan_status, d.status AS device_status,
               d.borrower_name AS device_borrower_name,
               d.borrower_user_id AS device_borrower_user_id,
               d.expected_return_at AS device_expected_return_at,
               rq.requested_at
        FROM borrow_records br
        JOIN devices d ON d.id = br.device_id
        LEFT JOIN borrow_requests rq ON rq.id = br.request_id
        WHERE br.id = ?
        """,
        (record_id,),
    ).fetchone()
    return dict(row) if row else {}


def _validate_manual_overdue_context(item: Dict[str, Any], now: datetime) -> None:
    if not item:
        raise HTTPException(status_code=404, detail="借用记录不存在")
    if item.get("record_status") != "borrowed" or item.get("returned_at") is not None:
        raise HTTPException(status_code=400, detail="借用记录已归还，不能发送逾期通知")
    if not item.get("expected_return_at"):
        raise HTTPException(status_code=400, detail="借用记录缺少预计归还时间")
    expected = _parse_datetime(item["expected_return_at"])
    if expected >= now:
        raise HTTPException(status_code=400, detail="借用记录尚未超过预计归还时间")
    if not _clean_text(item.get("borrower_user_id")):
        raise HTTPException(status_code=400, detail="借用人缺少门户用户 ID，无法发送通知")
    if item.get("device_loan_status") != "borrowed":
        raise HTTPException(status_code=409, detail="当前设备已不处于借用中")
    if item.get("device_status") == RESIDENT_DEVICE_STATUS:
        raise HTTPException(status_code=409, detail="当前设备状态不适合发送逾期通知")


def _upsert_manual_overdue_notification(conn, item: Dict[str, Any], now: datetime) -> Dict[str, Any]:
    _validate_manual_overdue_context(item, now)
    current = now_iso()
    borrower_user_id = _clean_text(item["borrower_user_id"])
    borrower_name = _clean_text(item["borrower_name"])
    requested_at = item.get("requested_at") or item.get("borrowed_at")
    conn.execute(
        """
        INSERT OR IGNORE INTO overdue_notifications (
            device_id, borrow_record_id, borrower_user_id, borrower_name, device_model,
            requested_at, expected_return_at, status, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, 'pending', ?, ?)
        """,
        (
            item["device_id"],
            item["borrow_record_id"],
            borrower_user_id,
            borrower_name,
            item["device_model"],
            requested_at,
            item["expected_return_at"],
            current,
            current,
        ),
    )
    row = conn.execute(
        """
        SELECT id, device_id, borrow_record_id, borrower_user_id, borrower_name,
               device_model, requested_at, expected_return_at
        FROM overdue_notifications
        WHERE borrow_record_id = ? AND borrower_user_id = ? AND expected_return_at = ?
        ORDER BY id DESC
        LIMIT 1
        """,
        (item["borrow_record_id"], borrower_user_id, item["expected_return_at"]),
    ).fetchone()
    notification = dict(row) if row else {}
    if not notification or not _is_overdue_snapshot_current(conn, notification, now):
        raise HTTPException(status_code=409, detail="借用记录状态已变化，无法发送逾期通知")
    return notification


def _record_manual_overdue_notification_error(conn, notification_id: int, exc: PortalNotificationError) -> None:
    status_code, code, message = _portal_error_parts(exc)
    conn.execute(
        """
        UPDATE overdue_notifications
        SET manual_last_error_code = ?, manual_last_error_message = ?,
            manual_last_error_status = ?, updated_at = ?
        WHERE id = ?
        """,
        (code, message, status_code, now_iso(), notification_id),
    )


def _mark_manual_overdue_notification_sent(conn, notification: Dict[str, Any], sent_at: str) -> bool:
    now = _parse_datetime(sent_at)
    if not _is_overdue_snapshot_current(conn, notification, now):
        _mark_overdue_notification_skipped(
            conn,
            notification["id"],
            "OVERDUE_NOTIFICATION_STALE",
            "逾期通知快照已失效",
        )
        return False
    cur = conn.execute(
        """
        UPDATE overdue_notifications
        SET status = 'sent',
            sent_at = COALESCE(sent_at, ?),
            manual_sent_at = ?,
            manual_last_error_code = NULL,
            manual_last_error_message = NULL,
            manual_last_error_status = NULL,
            updated_at = ?
        WHERE id = ?
        """,
        (sent_at, sent_at, sent_at, notification["id"]),
    )
    if cur.rowcount == 0:
        return False
    conn.execute(
        "UPDATE devices SET overdue_notified = 1, updated_at = ? WHERE id = ?",
        (sent_at, notification["device_id"]),
    )
    return True


async def _send_overdue_webhook_notifications(items: List[Dict[str, Any]]) -> None:
    with db_session() as conn:
        webhook = _get_setting(conn, "feishu_webhook")
        defaults = _build_configured_webhook_message(
            conn,
            trigger="overdue",
            fallback_title="逾期通知",
            fallback_body="",
        )
    if not webhook:
        return
    client: httpx.AsyncClient = app.state.http_client
    for item in items:
        fallback_body = (
            f"借用人名字: {item['borrower_name'] or '-'}\n"
            f"借用的设备型号: {item['device_model']}\n"
            f"预计归还时间: {_format_notify_time(item['expected_return_at'])}\n"
            "借用时间已逾期"
        )
        template_values = {
            "borrower": item["borrower_name"] or "-",
            "device_model": item["device_model"],
            "return_time": _format_notify_time(item["expected_return_at"]),
        }
        title = _render_notification_template(defaults["title"], template_values)
        body = _render_notification_template(defaults["body"] or fallback_body, template_values)
        try:
            await send_feishu_message(
                client,
                webhook,
                title,
                body,
                card_color=defaults["card_color"],
                admin_url=defaults["admin_url"],
            )
        except Exception as exc:
            logger.warning("旧飞书逾期通知发送失败: notification_id=%s", item["id"])
            with db_session() as conn:
                _record_overdue_webhook_error(conn, item["id"], str(exc))
            continue
        sent_at = now_iso()
        with db_session() as conn:
            if _is_overdue_snapshot_current(conn, item, _parse_datetime(sent_at), require_recipient=False):
                _mark_overdue_webhook_sent(conn, item, sent_at)


async def _process_overdue_notifications_once(now: Optional[datetime] = None) -> None:
    checked_at = now or datetime.now(timezone.utc)
    with db_session() as conn:
        _upsert_pending_overdue_notifications(conn, checked_at)
        webhook_pending = _pending_overdue_webhook_notifications(conn, checked_at)
        pending = _pending_overdue_notifications(conn, checked_at)
        auth = _portal_notification_auth_from_job()
        portal_auth_missing = pending and not auth.get("service_token")
        if portal_auth_missing:
            _record_overdue_auth_missing(conn)

    await _send_overdue_webhook_notifications(webhook_pending)
    if portal_auth_missing:
        return
    if not pending:
        return

    client: httpx.AsyncClient = app.state.http_client
    for item in pending:
        payload = _build_configured_portal_card_payload(
            trigger="overdue",
            borrower=item["borrower_name"],
            device_name=item["device_model"],
            request_date=item["requested_at"],
            return_date=item["expected_return_at"],
        )
        try:
            await send_portal_notification(
                client,
                recipient_user_id=item["borrower_user_id"] or "",
                payload=payload,
                service_id=auth.get("service_id"),
                service_token=auth.get("service_token"),
            )
        except PortalNotificationError as exc:
            _log_portal_notification_error("overdue", exc)
            with db_session() as conn:
                _record_overdue_notification_error(conn, item["id"], exc)
            continue
        sent_at = now_iso()
        with db_session() as conn:
            _mark_overdue_notification_sent(conn, item, sent_at)


def _overdue_notification_status_summary(conn) -> Dict[str, Any]:
    rows = conn.execute(
        """
        SELECT status, COUNT(1) AS count
        FROM overdue_notifications
        GROUP BY status
        """
    ).fetchall()
    recent = conn.execute(
        """
        SELECT id, device_id, borrower_name, device_model, expected_return_at, status,
               last_error_code, last_error_message, last_error_status,
               webhook_sent_at, webhook_last_error_message, sent_at, updated_at
        FROM overdue_notifications
        ORDER BY updated_at DESC
        LIMIT 20
        """
    ).fetchall()
    return {
        "items_by_status": {row["status"]: row["count"] for row in rows},
        "recent_items": [dict(row) for row in recent],
    }


async def _queue_notify(
    title: str,
    body: str,
    trigger: Optional[str] = None,
    template_values: Optional[Dict[str, Any]] = None,
):
    webhook = None
    with db_session() as conn:
        webhook = _get_setting(conn, "feishu_webhook")
        webhook_message = _build_configured_webhook_message(
            conn,
            trigger=trigger,
            fallback_title=title,
            fallback_body=body,
            template_values=template_values,
        )
    if not webhook:
        return
    client: httpx.AsyncClient = app.state.http_client
    try:
        await send_feishu_message(
            client,
            webhook,
            webhook_message["title"],
            webhook_message["body"],
            card_color=webhook_message["card_color"],
            admin_url=webhook_message["admin_url"],
        )
    except NotifyError:
        # 避免通知失败影响主流程
        return
    except Exception:
        logger.warning("旧飞书通知发送失败: title=%s", title)
        return


async def _overdue_worker():
    while True:
        try:
            await asyncio.sleep(60)
            await _process_overdue_notifications_once()
        except asyncio.CancelledError:
            break
        except Exception:
            logger.warning("逾期通知任务执行失败", exc_info=True)
            continue


@app.on_event("startup")
async def on_startup():
    init_db()
    limits = None
    if hasattr(httpx, "Limits"):
        limits = httpx.Limits(max_connections=20, max_keepalive_connections=10)
    if limits:
        app.state.http_client = httpx.AsyncClient(timeout=20.0, limits=limits)
    else:
        app.state.http_client = httpx.AsyncClient(timeout=20.0)
    app.state.overdue_task = asyncio.create_task(_overdue_worker())


@app.on_event("shutdown")
async def on_shutdown():
    task = app.state.overdue_task
    task.cancel()
    try:
        await task
    except asyncio.CancelledError:
        pass
    await app.state.http_client.aclose()


@app.get("/api/vendors")
async def list_vendors():
    with db_session() as conn:
        rows = conn.execute("SELECT id, name FROM vendors ORDER BY id DESC").fetchall()
    return {"items": [dict(row) for row in rows]}


@app.post("/api/vendors")
async def create_vendor(payload: VendorCreate):
    _ensure_required(payload.name, "厂商")
    with db_session() as conn:
        try:
            conn.execute("INSERT INTO vendors (name) VALUES (?)", (payload.name.strip(),))
        except Exception:
            raise HTTPException(status_code=400, detail="厂商已存在")
    return {"message": "新增成功"}


@app.put("/api/vendors/{vendor_id}")
async def update_vendor(vendor_id: int, payload: VendorUpdate):
    _ensure_required(payload.name, "厂商")
    with db_session() as conn:
        try:
            cur = conn.execute("UPDATE vendors SET name = ? WHERE id = ?", (payload.name.strip(), vendor_id))
        except Exception:
            raise HTTPException(status_code=400, detail="厂商名称重复")
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="厂商不存在")
    return {"message": "更新成功"}


@app.delete("/api/vendors/{vendor_id}")
async def delete_vendor(vendor_id: int, payload: VendorDeleteRequest):
    with db_session() as conn:
        existing = conn.execute("SELECT id FROM vendors WHERE id = ?", (vendor_id,)).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="厂商不存在")
        device_count = conn.execute(
            "SELECT COUNT(1) AS cnt FROM devices WHERE vendor_id = ?",
            (vendor_id,),
        ).fetchone()["cnt"]
        if device_count > 0 and not payload.rebind_vendor_id:
            raise HTTPException(status_code=409, detail="厂商已绑定设备，需要重新绑定")
        if device_count > 0:
            if payload.rebind_vendor_id == vendor_id:
                raise HTTPException(status_code=400, detail="新的厂商不能与原厂商相同")
            target = conn.execute(
                "SELECT id FROM vendors WHERE id = ?",
                (payload.rebind_vendor_id,),
            ).fetchone()
            if not target:
                raise HTTPException(status_code=400, detail="新的厂商不存在")
            conn.execute(
                "UPDATE devices SET vendor_id = ? WHERE vendor_id = ?",
                (payload.rebind_vendor_id, vendor_id),
            )
        conn.execute("DELETE FROM vendors WHERE id = ?", (vendor_id,))
    return {"message": "删除成功"}


@app.get("/api/systems")
async def list_systems(include_versions: bool = Query(False)):
    with db_session() as conn:
        systems = conn.execute("SELECT id, name FROM systems ORDER BY id DESC").fetchall()
        system_list = [dict(row) for row in systems]
        if include_versions:
            versions = conn.execute(
                "SELECT id, system_id, version FROM system_versions ORDER BY id DESC"
            ).fetchall()
            versions_by_system: Dict[int, List[Dict[str, Any]]] = {}
            for row in versions:
                versions_by_system.setdefault(row["system_id"], []).append(dict(row))
            for item in system_list:
                item["versions"] = versions_by_system.get(item["id"], [])
    return {"items": system_list}


@app.post("/api/systems")
async def create_system(payload: SystemCreate):
    _ensure_required(payload.name, "系统")
    with db_session() as conn:
        try:
            conn.execute("INSERT INTO systems (name) VALUES (?)", (payload.name.strip(),))
        except Exception:
            raise HTTPException(status_code=400, detail="系统已存在")
    return {"message": "新增成功"}


@app.put("/api/systems/{system_id}")
async def update_system(system_id: int, payload: SystemUpdate):
    _ensure_required(payload.name, "系统")
    with db_session() as conn:
        try:
            cur = conn.execute("UPDATE systems SET name = ? WHERE id = ?", (payload.name.strip(), system_id))
        except Exception:
            raise HTTPException(status_code=400, detail="系统名称重复")
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="系统不存在")
    return {"message": "更新成功"}


@app.delete("/api/systems/{system_id}")
async def delete_system(system_id: int, payload: SystemDeleteRequest):
    with db_session() as conn:
        existing = conn.execute("SELECT id FROM systems WHERE id = ?", (system_id,)).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="系统不存在")
        device_count = conn.execute(
            "SELECT COUNT(1) AS cnt FROM devices WHERE system_id = ?",
            (system_id,),
        ).fetchone()["cnt"]
        if device_count > 0 and (not payload.rebind_system_id or not payload.rebind_version_id):
            raise HTTPException(status_code=409, detail="系统已绑定设备，需要重新绑定")
        if device_count > 0:
            if payload.rebind_system_id == system_id:
                raise HTTPException(status_code=400, detail="新的系统不能与原系统相同")
            target = conn.execute(
                "SELECT id FROM systems WHERE id = ?",
                (payload.rebind_system_id,),
            ).fetchone()
            if not target:
                raise HTTPException(status_code=400, detail="新的系统不存在")
            _validate_version_for_system(conn, payload.rebind_system_id, payload.rebind_version_id)
            conn.execute(
                "UPDATE devices SET system_id = ?, system_version_id = ? WHERE system_id = ?",
                (payload.rebind_system_id, payload.rebind_version_id, system_id),
            )
        conn.execute("DELETE FROM systems WHERE id = ?", (system_id,))
    return {"message": "删除成功"}


@app.post("/api/systems/{system_id}/versions")
async def create_version(system_id: int, payload: VersionCreate):
    _ensure_required(payload.version, "版本")
    with db_session() as conn:
        system = conn.execute("SELECT id FROM systems WHERE id = ?", (system_id,)).fetchone()
        if not system:
            raise HTTPException(status_code=404, detail="系统不存在")
        try:
            conn.execute(
                "INSERT INTO system_versions (system_id, version) VALUES (?, ?)",
                (system_id, payload.version.strip()),
            )
        except Exception:
            raise HTTPException(status_code=400, detail="版本已存在")
    return {"message": "新增成功"}


@app.put("/api/versions/{version_id}")
async def update_version(version_id: int, payload: VersionUpdate):
    _ensure_required(payload.version, "版本")
    with db_session() as conn:
        try:
            cur = conn.execute(
                "UPDATE system_versions SET version = ? WHERE id = ?",
                (payload.version.strip(), version_id),
            )
        except Exception:
            raise HTTPException(status_code=400, detail="版本名称重复")
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="版本不存在")
    return {"message": "更新成功"}


@app.delete("/api/versions/{version_id}")
async def delete_version(version_id: int, payload: VersionDeleteRequest):
    with db_session() as conn:
        existing = conn.execute(
            "SELECT id, system_id FROM system_versions WHERE id = ?", (version_id,)
        ).fetchone()
        if not existing:
            raise HTTPException(status_code=404, detail="版本不存在")
        device_count = conn.execute(
            "SELECT COUNT(1) AS cnt FROM devices WHERE system_version_id = ?",
            (version_id,),
        ).fetchone()["cnt"]
        if device_count > 0 and not payload.rebind_version_id:
            raise HTTPException(status_code=409, detail="版本已绑定设备，需要重新绑定")
        if device_count > 0:
            if payload.rebind_version_id == version_id:
                raise HTTPException(status_code=400, detail="新的版本不能与原版本相同")
            target = conn.execute(
                "SELECT id, system_id FROM system_versions WHERE id = ?",
                (payload.rebind_version_id,),
            ).fetchone()
            if not target:
                raise HTTPException(status_code=400, detail="新的版本不存在")
            if target["system_id"] != existing["system_id"]:
                raise HTTPException(status_code=400, detail="新版本必须属于同一系统")
            conn.execute(
                "UPDATE devices SET system_version_id = ? WHERE system_version_id = ?",
                (payload.rebind_version_id, version_id),
            )
        conn.execute("DELETE FROM system_versions WHERE id = ?", (version_id,))
    return {"message": "删除成功"}


@app.get("/api/devices")
async def list_devices(query: Optional[str] = None):
    with db_session() as conn:
        items = _fetch_devices(conn, query)
    return {"items": items}


@app.post("/api/current-user/migrate-borrower-data")
async def migrate_current_user_borrower_data(request: Request):
    profile = await _require_portal_borrower_profile(request)
    with db_session() as conn:
        migrated = 0
        for table in ("devices", "borrow_requests", "borrow_records"):
            migrated += _migrate_borrower_table(conn, table, profile)
        migrated += _migrate_borrower_change_side(conn, "borrower_before", profile)
        migrated += _migrate_borrower_change_side(conn, "borrower_after", profile)
    return {"message": "迁移完成", "migrated": migrated}


@app.post("/api/devices")
async def create_device(payload: DeviceCreate):
    _ensure_required(payload.model, "设备型号")
    _ensure_required(payload.status, "设备状态")
    with db_session() as conn:
        vendor = conn.execute("SELECT id FROM vendors WHERE id = ?", (payload.vendor_id,)).fetchone()
        if not vendor:
            raise HTTPException(status_code=400, detail="厂商不存在")
        system = conn.execute("SELECT id FROM systems WHERE id = ?", (payload.system_id,)).fetchone()
        if not system:
            raise HTTPException(status_code=400, detail="系统不存在")
        _validate_version_for_system(conn, payload.system_id, payload.system_version_id)
        now = now_iso()
        conn.execute(
            """
            INSERT INTO devices (
                model, status, type, vendor_id, system_id, system_version_id,
                resolution, arch, cpu, boot_password, notes,
                loan_status, borrower_name, borrowed_at, expected_return_at,
                overdue_notified, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'available', NULL, NULL, NULL, 0, ?, ?)
            """,
            (
                payload.model.strip(),
                payload.status.strip(),
                payload.type,
                payload.vendor_id,
                payload.system_id,
                payload.system_version_id,
                payload.resolution,
                payload.arch,
                payload.cpu,
                payload.boot_password,
                payload.notes,
                now,
                now,
            ),
        )
    return {"message": "新增成功"}


@app.put("/api/devices/{device_id}")
async def update_device(device_id: int, payload: DeviceUpdate):
    _ensure_required(payload.model, "设备型号")
    _ensure_required(payload.status, "设备状态")
    with db_session() as conn:
        vendor = conn.execute("SELECT id FROM vendors WHERE id = ?", (payload.vendor_id,)).fetchone()
        if not vendor:
            raise HTTPException(status_code=400, detail="厂商不存在")
        system = conn.execute("SELECT id FROM systems WHERE id = ?", (payload.system_id,)).fetchone()
        if not system:
            raise HTTPException(status_code=400, detail="系统不存在")
        _validate_version_for_system(conn, payload.system_id, payload.system_version_id)
        now = now_iso()
        cur = conn.execute(
            """
            UPDATE devices SET
                model = ?, status = ?, type = ?, vendor_id = ?, system_id = ?, system_version_id = ?,
                resolution = ?, arch = ?, cpu = ?, boot_password = ?, notes = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                payload.model.strip(),
                payload.status.strip(),
                payload.type,
                payload.vendor_id,
                payload.system_id,
                payload.system_version_id,
                payload.resolution,
                payload.arch,
                payload.cpu,
                payload.boot_password,
                payload.notes,
                now,
                device_id,
            ),
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="设备不存在")
    return {"message": "更新成功"}


@app.delete("/api/devices/{device_id}")
async def delete_device(device_id: int):
    with db_session() as conn:
        cur = conn.execute("DELETE FROM devices WHERE id = ?", (device_id,))
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="设备不存在")
    return {"message": "删除成功"}


@app.post("/api/devices/{device_id}/borrow")
async def borrow_device(device_id: int, payload: BorrowRequest, background_tasks: BackgroundTasks, request: Request):
    borrower_profile = await _resolve_borrower_profile(request, payload.borrower_name)
    portal_auth = _portal_notification_auth_from_request(request)
    _ensure_required(borrower_profile["name"], "借用人名字")
    expected_return = _parse_datetime(payload.expected_return_at)
    now = datetime.now(timezone.utc)
    if expected_return <= now:
        raise HTTPException(status_code=400, detail="预计归还时间必须晚于当前时间")
    with db_session() as conn:
        row = conn.execute(
            "SELECT loan_status, status, model FROM devices WHERE id = ?",
            (device_id,),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="设备不存在")
        if row["loan_status"] != "available":
            raise HTTPException(status_code=400, detail="设备不可借用")
        if row["status"] == "未登记借用":
            raise HTTPException(
                status_code=400,
                detail="未找到该设备的借用人，无法进行设备借用，请找回设备后，把状态改回“正常”。",
            )
        request_now = now_iso()
        cur = conn.execute(
            """
            INSERT INTO borrow_requests (
                device_id, device_model, borrower_name, borrower_user_id, borrower_open_id,
                borrower_avatar_url, borrower_job_title, expected_return_at,
                request_type, status, requested_at, handled_at, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'borrow', 'pending', ?, NULL, ?, ?)
            """,
            (
                device_id,
                row["model"],
                *_borrower_values(borrower_profile),
                expected_return.isoformat(),
                request_now,
                request_now,
                request_now,
            ),
        )
        conn.execute(
            """
            UPDATE devices SET
                loan_status = 'pending',
                borrower_name = ?,
                borrower_user_id = ?,
                borrower_open_id = ?,
                borrower_avatar_url = ?,
                borrower_job_title = ?,
                borrowed_at = NULL,
                expected_return_at = ?,
                overdue_notified = 0,
                updated_at = ?
            WHERE id = ?
            """,
            (
                *_borrower_values(borrower_profile),
                expected_return.isoformat(),
                now_iso(),
                device_id,
            ),
        )
    body = (
        f"借用人名字: {borrower_profile['name']}\n"
        f"借用的设备型号: {row['model']}\n"
        f"归还时间: {_format_notify_time(expected_return)}\n"
        "待借设备"
    )
    background_tasks.add_task(
        _queue_notify,
        "待借通知",
        body,
        "borrow_submit",
        {
            "borrower": borrower_profile["name"],
            "device_model": row["model"],
            "return_time": _format_notify_time(expected_return),
        },
    )
    _add_portal_notification_task(
        background_tasks,
        context="borrow_submit",
        recipient_user_id=borrower_profile.get("user_id"),
        payload=_build_configured_portal_card_payload(
            trigger="borrow_submit",
            borrower=borrower_profile["name"],
            device_name=row["model"],
            request_date=request_now,
            return_date=expected_return,
        ),
        auth=portal_auth,
    )
    return {"message": "已提交待借申请", "request_id": cur.lastrowid}


@app.post("/api/devices/{device_id}/extend")
async def extend_device(device_id: int, payload: ExtendRequest, background_tasks: BackgroundTasks, request: Request):
    portal_auth = _portal_notification_auth_from_request(request)
    expected_return = _parse_datetime(payload.expected_return_at)
    with db_session() as conn:
        row = conn.execute(
            "SELECT loan_status, expected_return_at, borrower_name, borrower_user_id, model FROM devices WHERE id = ?",
            (device_id,),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="设备不存在")
        if row["loan_status"] != "borrowed":
            raise HTTPException(status_code=400, detail="设备未借出")
        old_expected = _parse_datetime(row["expected_return_at"]) if row["expected_return_at"] else None
        if old_expected and expected_return <= old_expected:
            raise HTTPException(status_code=400, detail="延期时间必须晚于当前预计归还时间")
        borrow_context = _fetch_current_borrow_context(conn, device_id)
        conn.execute(
            """
            UPDATE devices SET expected_return_at = ?, overdue_notified = 0, updated_at = ?
            WHERE id = ?
            """,
            (expected_return.isoformat(), now_iso(), device_id),
        )
        _update_borrow_record_expected(conn, device_id, expected_return.isoformat())
    body = (
        f"借用人名字: {row['borrower_name'] or '-'}\n"
        f"借用的设备型号: {row['model']}\n"
        f"预计归还时间(旧): {_format_notify_time(row['expected_return_at'])}\n"
        f"变更为 预计归还时间(新): {_format_notify_time(expected_return)}"
    )
    background_tasks.add_task(
        _queue_notify,
        "延期通知",
        body,
        "extend",
        {
            "borrower": row["borrower_name"] or "-",
            "device_model": row["model"],
            "old_return_time": _format_notify_time(row["expected_return_at"]),
            "return_time": _format_notify_time(expected_return),
            "new_return_time": _format_notify_time(expected_return),
        },
    )
    _add_portal_notification_task(
        background_tasks,
        context="extend",
        recipient_user_id=row["borrower_user_id"],
        payload=_build_configured_portal_card_payload(
            trigger="extend",
            borrower=row["borrower_name"] or "-",
            device_name=row["model"],
            request_date=_borrow_context_request_date(borrow_context),
            return_date=expected_return,
        ),
        auth=portal_auth,
    )
    return {"message": "延期成功"}


@app.post("/api/devices/{device_id}/change-borrower")
async def change_borrower(
    device_id: int, payload: BorrowerChangeRequest, background_tasks: BackgroundTasks, request: Request
):
    borrower_profile = await _resolve_borrower_profile(request, payload.borrower_name)
    portal_auth = _portal_notification_auth_from_request(request)
    _ensure_required(borrower_profile["name"], "借用人名字")
    expected_return = _parse_datetime(payload.expected_return_at)
    now = datetime.now(timezone.utc)
    if expected_return <= now:
        raise HTTPException(status_code=400, detail="预计归还时间必须晚于当前时间")
    with db_session() as conn:
        row = conn.execute(
            """
            SELECT loan_status, borrower_name, borrower_user_id, borrower_open_id,
                   borrower_avatar_url, borrower_job_title, expected_return_at, model
            FROM devices WHERE id = ?
            """,
            (device_id,),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="设备不存在")
        if row["loan_status"] not in {"borrowed", "pending"}:
            raise HTTPException(status_code=400, detail="设备未处于借用状态")
        if not row["borrower_name"]:
            raise HTTPException(status_code=400, detail="当前设备无借用人")
        pending_change = conn.execute(
            """
            SELECT id FROM borrow_requests
            WHERE device_id = ? AND request_type = 'change' AND status = 'pending'
            LIMIT 1
            """,
            (device_id,),
        ).fetchone()
        if pending_change:
            raise HTTPException(status_code=400, detail="当前设备已有借用人更换申请，需等待管理员处理。")
        old_borrower = row["borrower_name"]
        old_expected = row["expected_return_at"]
        old_borrow_context = _fetch_current_borrow_context(conn, device_id)
        request_now = now_iso()
        conn.execute(
            """
            INSERT INTO borrow_requests (
                device_id, device_model, borrower_name, borrower_user_id, borrower_open_id,
                borrower_avatar_url, borrower_job_title, expected_return_at,
                request_type, status, requested_at, handled_at, created_at, updated_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'change', 'pending', ?, NULL, ?, ?)
            """,
            (
                device_id,
                row["model"],
                *_borrower_values(borrower_profile),
                expected_return.isoformat(),
                request_now,
                request_now,
                request_now,
            ),
        )
    body = (
        f"设备: {row['model']}\n"
        f"变更前借用人: {old_borrower or '-'}\n"
        f"变更前归还时间: {_format_notify_time(old_expected)}\n"
        f"变更后借用人: {borrower_profile['name']}\n"
        f"变更后预期归还时间: {_format_notify_time(expected_return)}\n"
        f"变更时间: {_format_notify_time(now)}"
    )
    background_tasks.add_task(
        _queue_notify,
        "借用人变更通知",
        body,
        "change_borrower_submit_old",
        {
            "borrower": borrower_profile["name"],
            "device_model": row["model"],
            "old_borrower": old_borrower or "-",
            "new_borrower": borrower_profile["name"],
            "old_return_time": _format_notify_time(old_expected),
            "return_time": _format_notify_time(expected_return),
            "new_return_time": _format_notify_time(expected_return),
            "change_time": _format_notify_time(now),
        },
    )
    _add_portal_notification_task(
        background_tasks,
        context="change_borrower_submit_old",
        recipient_user_id=row["borrower_user_id"],
        payload=_build_configured_portal_card_payload(
            trigger="change_borrower_submit_old",
            borrower=old_borrower,
            device_name=row["model"],
            request_date=_borrow_context_request_date(old_borrow_context),
            return_date=old_expected,
            template_values={"old_borrower": old_borrower, "new_borrower": borrower_profile["name"]},
        ),
        auth=portal_auth,
    )
    _add_portal_notification_task(
        background_tasks,
        context="change_borrower_submit_new",
        recipient_user_id=borrower_profile.get("user_id"),
        payload=_build_configured_portal_card_payload(
            trigger="change_borrower_submit_new",
            borrower=borrower_profile["name"],
            device_name=row["model"],
            request_date=request_now,
            return_date=expected_return,
            template_values={"old_borrower": old_borrower, "new_borrower": borrower_profile["name"]},
        ),
        auth=portal_auth,
    )
    return {"message": "已提交借用人变更申请"}


@app.post("/api/devices/{device_id}/return")
async def return_device(device_id: int, background_tasks: BackgroundTasks, request: Request):
    portal_auth = _portal_notification_auth_from_request(request)
    now = datetime.now(timezone.utc)
    with db_session() as conn:
        row = conn.execute(
            "SELECT loan_status, borrower_name, borrower_user_id, model FROM devices WHERE id = ?",
            (device_id,),
        ).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="设备不存在")
        if row["loan_status"] != "borrowed":
            raise HTTPException(status_code=400, detail="设备未处于已借出状态")
        borrow_context = _fetch_current_borrow_context(conn, device_id)
        conn.execute(
            """
            UPDATE devices SET
                loan_status = 'available',
                borrower_name = NULL,
                borrower_user_id = NULL,
                borrower_open_id = NULL,
                borrower_avatar_url = NULL,
                borrower_job_title = NULL,
                borrowed_at = NULL,
                expected_return_at = NULL,
                overdue_notified = 0,
                updated_at = ?
            WHERE id = ?
            """,
            (now_iso(), device_id),
        )
        _close_borrow_record(conn, device_id, now_iso())
    body = (
        f"借用人名字: {row['borrower_name'] or '-'}\n"
        f"借用的设备型号: {row['model']}\n"
        f"归还时间: {_format_notify_time(now)}\n"
        "归还成功"
    )
    background_tasks.add_task(
        _queue_notify,
        "归还通知",
        body,
        "return",
        {
            "borrower": row["borrower_name"] or "-",
            "device_model": row["model"],
            "return_time": _format_notify_time(now),
        },
    )
    _add_portal_notification_task(
        background_tasks,
        context="return",
        recipient_user_id=row["borrower_user_id"],
        payload=_build_configured_portal_card_payload(
            trigger="return",
            borrower=row["borrower_name"] or "-",
            device_name=row["model"],
            request_date=_borrow_context_request_date(borrow_context),
            return_date=now,
        ),
        auth=portal_auth,
    )
    return {"message": "归还成功"}


@app.get("/api/borrow-requests")
async def list_borrow_requests(
    query: Optional[str] = Query(default=None),
    status: Optional[str] = Query(default=None),
    device_id: Optional[int] = Query(default=None),
    request_type: Optional[str] = Query(default=None),
):
    if status and status not in {"pending", "approved", "cancelled"}:
        raise HTTPException(status_code=400, detail="状态参数不合法")
    with db_session() as conn:
        items = _fetch_borrow_requests(conn, query, status, device_id, request_type)
    return {"items": items}


@app.post("/api/borrow-requests/{request_id}/approve")
async def approve_borrow_request(request_id: int, background_tasks: BackgroundTasks, request_context: Request):
    portal_auth = _portal_notification_auth_from_request(request_context)
    now = datetime.now(timezone.utc)
    old_borrower_profile: Dict[str, Optional[str]] = {}
    new_borrower_profile: Dict[str, Optional[str]] = {}
    old_borrow_context: Dict[str, Any] = {}
    with db_session() as conn:
        request_row = conn.execute(
            "SELECT * FROM borrow_requests WHERE id = ?",
            (request_id,),
        ).fetchone()
        request = _row_to_dict(request_row)
        if not request:
            raise HTTPException(status_code=404, detail="申请不存在")
        if request["status"] != "pending":
            raise HTTPException(status_code=400, detail="申请已处理")
        device = conn.execute(
            """
            SELECT loan_status, model, borrower_name, borrower_user_id, borrower_open_id,
                   borrower_avatar_url, borrower_job_title, expected_return_at
            FROM devices WHERE id = ?
            """,
            (request["device_id"],),
        ).fetchone()
        if not device:
            raise HTTPException(status_code=404, detail="设备不存在")

        if request["request_type"] == "change":
            if device["loan_status"] not in {"borrowed", "pending"}:
                raise HTTPException(status_code=400, detail="设备状态异常")
            old_expected = device["expected_return_at"]
            old_borrower_profile = _borrower_profile_from_row(device)
            new_borrower_profile = _borrower_profile_from_row(request)
            old_borrow_context = _fetch_current_borrow_context(conn, request["device_id"])
            conn.execute(
                """
                UPDATE devices SET
                    borrower_name = ?,
                    borrower_user_id = ?,
                    borrower_open_id = ?,
                    borrower_avatar_url = ?,
                    borrower_job_title = ?,
                    expected_return_at = ?,
                    overdue_notified = 0,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    *_borrower_values(new_borrower_profile),
                    request["expected_return_at"],
                    now_iso(),
                    request["device_id"],
                ),
            )
            record_row = conn.execute(
                """
                SELECT id FROM borrow_records
                WHERE device_id = ? AND returned_at IS NULL
                ORDER BY borrowed_at DESC
                LIMIT 1
                """,
                (request["device_id"],),
            ).fetchone()
            record_id = record_row["id"] if record_row else None
            change_request_id = request_id
            borrow_request_id = None
            if device["loan_status"] == "pending":
                borrow_request = conn.execute(
                    """
                    SELECT id FROM borrow_requests
                    WHERE device_id = ? AND request_type = 'borrow' AND status = 'pending'
                    ORDER BY id DESC
                    LIMIT 1
                    """,
                    (request["device_id"],),
                ).fetchone()
                if borrow_request:
                    borrow_request_id = borrow_request["id"]
                    conn.execute(
                        """
                        UPDATE borrow_requests
                        SET borrower_name = ?,
                            borrower_user_id = ?,
                            borrower_open_id = ?,
                            borrower_avatar_url = ?,
                            borrower_job_title = ?,
                            expected_return_at = ?,
                            updated_at = ?
                        WHERE id = ?
                        """,
                        (
                            *_borrower_values(new_borrower_profile),
                            request["expected_return_at"],
                            now_iso(),
                            borrow_request_id,
                        ),
                    )
            if record_id:
                conn.execute(
                    """
                    UPDATE borrow_records
                    SET borrower_name = ?,
                        borrower_user_id = ?,
                        borrower_open_id = ?,
                        borrower_avatar_url = ?,
                        borrower_job_title = ?,
                        expected_return_at = ?,
                        updated_at = ?
                    WHERE id = ?
                    """,
                    (
                        *_borrower_values(new_borrower_profile),
                        request["expected_return_at"],
                        now_iso(),
                        record_id,
                    ),
                )
            if record_id or borrow_request_id:
                now_str = now_iso()
                conn.execute(
                    """
                    INSERT INTO borrow_changes (
                        device_id, record_id, request_id,
                        borrower_before, borrower_before_user_id, borrower_before_open_id,
                        borrower_before_avatar_url, borrower_before_job_title,
                        borrower_after, borrower_after_user_id, borrower_after_open_id,
                        borrower_after_avatar_url, borrower_after_job_title,
                        expected_before, expected_after, changed_at, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        request["device_id"],
                        record_id,
                        borrow_request_id or change_request_id,
                        *_borrower_values(old_borrower_profile),
                        *_borrower_values(new_borrower_profile),
                        old_expected,
                        request["expected_return_at"],
                        now_str,
                        now_str,
                        now_str,
                    ),
                )
            conn.execute(
                """
                UPDATE borrow_requests
                SET status = 'approved', handled_at = ?, updated_at = ?
                WHERE id = ?
                """,
                (now_iso(), now_iso(), request_id),
            )
        else:
            if device["loan_status"] != "pending":
                raise HTTPException(status_code=400, detail="设备状态异常")
            borrowed_at = now_iso()
            conn.execute(
                """
                UPDATE devices SET
                    loan_status = 'borrowed',
                    borrower_name = ?,
                    borrower_user_id = ?,
                    borrower_open_id = ?,
                    borrower_avatar_url = ?,
                    borrower_job_title = ?,
                    borrowed_at = ?,
                    expected_return_at = ?,
                    overdue_notified = 0,
                    updated_at = ?
                WHERE id = ?
                """,
                (
                    *_borrower_values(_borrower_profile_from_row(request)),
                    borrowed_at,
                    request["expected_return_at"],
                    now_iso(),
                    request["device_id"],
                ),
            )
            conn.execute(
                """
                UPDATE borrow_requests
                SET status = 'approved', handled_at = ?, updated_at = ?
                WHERE id = ?
                """,
                (borrowed_at, now_iso(), request_id),
            )
            _insert_borrow_record(
                conn,
                request["device_id"],
                request["device_model"],
                _borrower_profile_from_row(request),
                borrowed_at,
                request["expected_return_at"],
                request_id,
            )

    if request["request_type"] == "change":
        body = (
            f"设备: {request['device_model']}\n"
            f"新的借用人名字: {request['borrower_name']}\n"
            f"新的归还时间: {_format_notify_time(request['expected_return_at'])}"
        )
        background_tasks.add_task(
            _queue_notify,
            "借用人变更成功通知",
            body,
            "change_borrower_approve_new",
            {
                "borrower": new_borrower_profile.get("name") or request["borrower_name"],
                "device_model": request["device_model"],
                "old_borrower": old_borrower_profile.get("name") or "-",
                "new_borrower": new_borrower_profile.get("name") or request["borrower_name"],
                "return_time": _format_notify_time(request["expected_return_at"]),
                "new_return_time": _format_notify_time(request["expected_return_at"]),
            },
        )
        _add_portal_notification_task(
            background_tasks,
            context="change_borrower_approve_old",
            recipient_user_id=old_borrower_profile.get("user_id"),
            payload=_build_configured_portal_card_payload(
                trigger="change_borrower_approve_old",
                borrower=old_borrower_profile.get("name"),
                device_name=request["device_model"],
                request_date=_borrow_context_request_date(old_borrow_context),
                return_date=old_borrow_context.get("expected_return_at") or old_expected,
                template_values={
                    "old_borrower": old_borrower_profile.get("name"),
                    "new_borrower": new_borrower_profile.get("name") or request["borrower_name"],
                },
            ),
            auth=portal_auth,
        )
        _add_portal_notification_task(
            background_tasks,
            context="change_borrower_approve_new",
            recipient_user_id=new_borrower_profile.get("user_id"),
            payload=_build_configured_portal_card_payload(
                trigger="change_borrower_approve_new",
                borrower=new_borrower_profile.get("name"),
                device_name=request["device_model"],
                request_date=request["requested_at"],
                return_date=request["expected_return_at"],
                template_values={
                    "old_borrower": old_borrower_profile.get("name"),
                    "new_borrower": new_borrower_profile.get("name") or request["borrower_name"],
                },
            ),
            auth=portal_auth,
        )
        return {"message": "借用人变更成功"}

    body = (
        f"借用人名字: {request['borrower_name']}\n"
        f"借用的设备型号: {request['device_model']}\n"
        f"借用时间: {_format_notify_time(now)}\n"
        f"预计归还时间: {_format_notify_time(request['expected_return_at'])}\n"
        "借用成功"
    )
    background_tasks.add_task(
        _queue_notify,
        "借用通知",
        body,
        "borrow_approve",
        {
            "borrower": request["borrower_name"],
            "device_model": request["device_model"],
            "borrow_time": _format_notify_time(now),
            "return_time": _format_notify_time(request["expected_return_at"]),
        },
    )
    _add_portal_notification_task(
        background_tasks,
        context="borrow_approve",
        recipient_user_id=request["borrower_user_id"],
        payload=_build_configured_portal_card_payload(
            trigger="borrow_approve",
            borrower=request["borrower_name"],
            device_name=request["device_model"],
            request_date=request["requested_at"],
            return_date=request["expected_return_at"],
        ),
        auth=portal_auth,
    )
    return {"message": "确认借出成功"}


@app.post("/api/borrow-requests/{request_id}/cancel")
async def cancel_borrow_request(request_id: int, background_tasks: BackgroundTasks, request_context: Request):
    portal_auth = _portal_notification_auth_from_request(request_context)
    now = now_iso()
    device_context: Dict[str, Any] = {}
    old_borrow_context: Dict[str, Any] = {}
    with db_session() as conn:
        request = conn.execute(
            "SELECT * FROM borrow_requests WHERE id = ?",
            (request_id,),
        ).fetchone()
        if not request:
            raise HTTPException(status_code=404, detail="申请不存在")
        if request["status"] != "pending":
            raise HTTPException(status_code=400, detail="申请已处理")
        device = conn.execute(
            """
            SELECT borrower_name, borrower_user_id, expected_return_at, model
            FROM devices WHERE id = ?
            """,
            (request["device_id"],),
        ).fetchone()
        device_context = dict(device) if device else {}
        old_borrow_context = _fetch_current_borrow_context(conn, request["device_id"])
        conn.execute(
            """
            UPDATE borrow_requests
            SET status = 'cancelled', handled_at = ?, updated_at = ?
            WHERE id = ?
            """,
            (now, now_iso(), request_id),
        )
        if request["request_type"] != "change":
            conn.execute(
                """
                UPDATE devices SET
                    loan_status = 'available',
                    borrower_name = NULL,
                    borrower_user_id = NULL,
                    borrower_open_id = NULL,
                    borrower_avatar_url = NULL,
                    borrower_job_title = NULL,
                    borrowed_at = NULL,
                    expected_return_at = NULL,
                    overdue_notified = 0,
                    updated_at = ?
                WHERE id = ?
                """,
                (now_iso(), request["device_id"]),
            )
    if request["request_type"] == "change":
        old_borrower_name = device_context.get("borrower_name") or "-"
        new_borrower_name = request["borrower_name"]
        body = (
            f"设备: {request['device_model']}\n"
            f"变更前借用人: {old_borrower_name}\n"
            f"变更后借用人: {new_borrower_name}\n"
            f"预计归还时间: {_format_notify_time(request['expected_return_at'])}\n"
            "借用人变更失败"
        )
        background_tasks.add_task(
            _queue_notify,
            "借用人变更失败通知",
            body,
            "change_borrower_cancel_new",
            {
                "borrower": new_borrower_name,
                "device_model": request["device_model"],
                "old_borrower": old_borrower_name,
                "new_borrower": new_borrower_name,
                "return_time": _format_notify_time(request["expected_return_at"]),
            },
        )
        _add_portal_notification_task(
            background_tasks,
            context="change_borrower_cancel_old",
            recipient_user_id=device_context.get("borrower_user_id"),
            payload=_build_configured_portal_card_payload(
                trigger="change_borrower_cancel_old",
                borrower=old_borrower_name,
                device_name=request["device_model"],
                request_date=_borrow_context_request_date(old_borrow_context),
                return_date=device_context.get("expected_return_at"),
            ),
            auth=portal_auth,
        )
        _add_portal_notification_task(
            background_tasks,
            context="change_borrower_cancel_new",
            recipient_user_id=request["borrower_user_id"],
            payload=_build_configured_portal_card_payload(
                trigger="change_borrower_cancel_new",
                borrower=new_borrower_name,
                device_name=request["device_model"],
                request_date=request["requested_at"],
                return_date=request["expected_return_at"],
            ),
            auth=portal_auth,
        )
    else:
        body = (
            f"借用人名字: {request['borrower_name']}\n"
            f"借用的设备型号: {request['device_model']}\n"
            f"归还时间: {_format_notify_time(request['expected_return_at'])}\n"
            "借用失败"
        )
        background_tasks.add_task(
            _queue_notify,
            "借用失败通知",
            body,
            "borrow_cancel",
            {
                "borrower": request["borrower_name"],
                "device_model": request["device_model"],
                "return_time": _format_notify_time(request["expected_return_at"]),
            },
        )
        _add_portal_notification_task(
            background_tasks,
            context="borrow_cancel",
            recipient_user_id=request["borrower_user_id"],
            payload=_build_configured_portal_card_payload(
                trigger="borrow_cancel",
                borrower=request["borrower_name"],
                device_name=request["device_model"],
                request_date=request["requested_at"],
                return_date=request["expected_return_at"],
            ),
            auth=portal_auth,
        )
    return {"message": "取消成功"}


@app.get("/api/borrow-records")
async def list_borrow_records(query: Optional[str] = Query(default=None)):
    with db_session() as conn:
        items = _fetch_borrow_records(conn, query)
    return {"items": items}


@app.post("/api/borrow-records/{record_id}/overdue-notification")
async def trigger_borrow_record_overdue_notification(record_id: int):
    auth = _portal_notification_auth_from_job()
    if not auth.get("service_token"):
        raise HTTPException(status_code=401, detail="缺少门户通知服务凭证，无法发送逾期通知")

    checked_at = datetime.now(timezone.utc)
    with db_session() as conn:
        item = _fetch_borrow_record_overdue_context(conn, record_id)
        notification = _upsert_manual_overdue_notification(conn, item, checked_at)

    payload = _build_configured_portal_card_payload(
        trigger="overdue",
        borrower=notification["borrower_name"],
        device_name=notification["device_model"],
        request_date=notification["requested_at"],
        return_date=notification["expected_return_at"],
    )
    client: httpx.AsyncClient = app.state.http_client
    try:
        await send_portal_notification(
            client,
            recipient_user_id=notification["borrower_user_id"] or "",
            payload=payload,
            service_id=auth.get("service_id"),
            service_token=auth.get("service_token"),
        )
    except PortalNotificationError as exc:
        _log_portal_notification_error("manual_overdue", exc)
        with db_session() as conn:
            _record_manual_overdue_notification_error(conn, notification["id"], exc)
        raise HTTPException(status_code=exc.status_code or 502, detail=_sanitize_error_message(exc.message))

    sent_at = now_iso()
    with db_session() as conn:
        if not _mark_manual_overdue_notification_sent(conn, notification, sent_at):
            raise HTTPException(status_code=409, detail="借用记录状态已变化，逾期通知发送状态未更新")
    return {"message": "逾期通知已发送", "manual_sent_at": sent_at}


@app.get("/api/devices/export")
async def export_devices(query: Optional[str] = Query(default=None)):
    with db_session() as conn:
        items = _fetch_devices(conn, query)
    wb = Workbook()
    ws = wb.active
    ws.title = "借用数据"
    ws.append(
        [
            "设备ID",
            "设备型号",
            "设备状态",
            "设备类型",
            "厂商",
            "系统",
            "系统版本",
            "分辨率",
            "架构",
            "CPU型号",
            "开机密码",
            "备注",
            "性能",
            "借用人",
            "借用时间",
            "预计归还时间",
            "借用状态",
        ]
    )
    for item in items:
        ws.append(
            [
                item.get("id") or "-",
                item.get("model") or "-",
                item.get("status") or "-",
                item.get("type") or "-",
                item.get("vendor_name") or "-",
                item.get("system_name") or "-",
                item.get("system_version") or "-",
                item.get("resolution") or "-",
                item.get("arch") or "-",
                item.get("cpu") or "-",
                item.get("boot_password") or "-",
                item.get("notes") or "-",
                _extract_performance(item.get("notes")),
                item.get("borrower_name") or "-",
                _format_notify_time(item.get("borrowed_at")),
                _format_notify_time(item.get("expected_return_at")),
                _format_loan_status(item.get("loan_status")),
            ]
        )
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = "borrow_data.xlsx"
    headers = {"Content-Disposition": f'attachment; filename="{filename}"'}
    return StreamingResponse(
        content=output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@app.get("/api/settings/feishu")
async def get_feishu_setting():
    with db_session() as conn:
        webhook = _get_setting(conn, "feishu_webhook")
        admin_url = _borrow_admin_url(conn)
    return {"webhook_url": webhook, "admin_url": admin_url}


@app.put("/api/settings/feishu")
async def update_feishu_setting(payload: SettingUpdate):
    _ensure_required(payload.webhook_url, "Webhook")
    with db_session() as conn:
        _set_setting(conn, "feishu_webhook", payload.webhook_url.strip())
        if payload.admin_url is not None:
            _ensure_required(payload.admin_url, "设备借用管理页")
            _set_setting(conn, BORROW_ADMIN_URL_SETTING_KEY, payload.admin_url.strip())
    return {"message": "保存成功"}


@app.get("/api/settings/notifications")
async def get_notification_settings():
    with db_session() as conn:
        overrides = _load_notification_param_overrides(conn)
    return {"items": _notification_settings_items(overrides), "color_options": list(NOTIFICATION_COLOR_OPTIONS)}


@app.put("/api/settings/notifications")
async def update_notification_settings(payload: NotificationSettingsUpdate):
    with db_session() as conn:
        overrides = _save_notification_settings(conn, payload.settings)
    return {
        "message": "保存成功",
        "items": _notification_settings_items(overrides),
        "color_options": list(NOTIFICATION_COLOR_OPTIONS),
    }


@app.get("/api/settings/webhook-notifications")
async def get_webhook_notification_settings():
    with db_session() as conn:
        overrides = _load_webhook_notification_param_overrides(conn)
        admin_url = _borrow_admin_url(conn)
    return {
        "items": _webhook_notification_settings_items(overrides),
        "color_options": list(NOTIFICATION_COLOR_OPTIONS),
        "admin_url": admin_url,
    }


@app.put("/api/settings/webhook-notifications")
async def update_webhook_notification_settings(payload: WebhookNotificationSettingsUpdate):
    with db_session() as conn:
        overrides = _save_webhook_notification_settings(conn, payload.settings)
        admin_url = _borrow_admin_url(conn)
    return {
        "message": "保存成功",
        "items": _webhook_notification_settings_items(overrides),
        "color_options": list(NOTIFICATION_COLOR_OPTIONS),
        "admin_url": admin_url,
    }


@app.get("/api/notifications/overdue-status")
async def get_overdue_notification_status():
    auth = _portal_notification_auth_from_job()
    with db_session() as conn:
        summary = _overdue_notification_status_summary(conn)
    return {
        "service_id": auth.get("service_id"),
        "service_token_configured": bool(os.environ.get(PORTAL_NOTIFICATION_SERVICE_TOKEN_ENV, "").strip()),
        "auth_source": auth.get("source"),
        **summary,
    }


@app.get("/api/llm/models")
async def list_llm_models():
    with db_session() as conn:
        rows = conn.execute(
            "SELECT id, name, api_type, base_url, api_key, model, max_tokens, is_default FROM llm_models ORDER BY id DESC"
        ).fetchall()
    return {"items": [dict(row) for row in rows]}


@app.get("/api/llm/models/assignments")
async def get_llm_assignments():
    with db_session() as conn:
        fast_id = _get_int_setting(conn, "llm_fast_model_id")
        accurate_id = _get_int_setting(conn, "llm_accurate_model_id")
    return {"fast_model_id": fast_id, "accurate_model_id": accurate_id}


def _reset_default(conn):
    conn.execute("UPDATE llm_models SET is_default = 0")


def _ensure_default(conn):
    row = conn.execute("SELECT id FROM llm_models WHERE is_default = 1 ORDER BY id LIMIT 1").fetchone()
    if not row:
        first = conn.execute("SELECT id FROM llm_models ORDER BY id LIMIT 1").fetchone()
        if first:
            conn.execute("UPDATE llm_models SET is_default = 1 WHERE id = ?", (first["id"],))


@app.post("/api/llm/models")
async def create_llm_model(payload: LLMModelCreate):
    with db_session() as conn:
        now = now_iso()
        if payload.is_default:
            _reset_default(conn)
        conn.execute(
            """
            INSERT INTO llm_models (name, api_type, base_url, api_key, model, max_tokens, is_default, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                payload.name.strip(),
                payload.api_type.strip(),
                payload.base_url.strip(),
                payload.api_key.strip(),
                payload.model.strip(),
                payload.max_tokens,
                1 if payload.is_default else 0,
                now,
                now,
            ),
        )
        _ensure_default(conn)
    return {"message": "新增成功"}


@app.post("/api/llm/models/{model_id}/assign")
async def assign_llm_model(model_id: int, payload: LLMModelAssignRequest):
    with db_session() as conn:
        _assign_llm_model(conn, model_id, payload.role)
    return {"message": "指派成功"}


@app.delete("/api/llm/models/assignments/{role}")
async def clear_llm_assignment(role: str):
    with db_session() as conn:
        _clear_llm_assignment(conn, role)
    return {"message": "取消指派成功"}


@app.put("/api/llm/models/{model_id}")
async def update_llm_model(model_id: int, payload: LLMModelUpdate):
    with db_session() as conn:
        if payload.is_default:
            _reset_default(conn)
        cur = conn.execute(
            """
            UPDATE llm_models SET
                name = ?, api_type = ?, base_url = ?, api_key = ?, model = ?, max_tokens = ?, is_default = ?, updated_at = ?
            WHERE id = ?
            """,
            (
                payload.name.strip(),
                payload.api_type.strip(),
                payload.base_url.strip(),
                payload.api_key.strip(),
                payload.model.strip(),
                payload.max_tokens,
                1 if payload.is_default else 0,
                now_iso(),
                model_id,
            ),
        )
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="模型不存在")
        _ensure_default(conn)
    return {"message": "更新成功"}


@app.delete("/api/llm/models/{model_id}")
async def delete_llm_model(model_id: int):
    with db_session() as conn:
        cur = conn.execute("DELETE FROM llm_models WHERE id = ?", (model_id,))
        if cur.rowcount == 0:
            raise HTTPException(status_code=404, detail="模型不存在")
        _ensure_default(conn)
    return {"message": "删除成功"}


@app.post("/api/llm/models/{model_id}/test")
async def test_llm_model(model_id: int):
    with db_session() as conn:
        row = conn.execute("SELECT * FROM llm_models WHERE id = ?", (model_id,)).fetchone()
        if not row:
            raise HTTPException(status_code=404, detail="模型不存在")
        config = dict(row)
    client: httpx.AsyncClient = app.state.http_client
    try:
        await call_llm(
            client,
            config,
            "你是一个测试助手。",
            "请返回 JSON: {\"device_ids\": []}",
        )
    except LLMError as exc:
        raise HTTPException(status_code=400, detail=f"模型不可用: {exc}")
    return {"message": "模型可用"}


@app.post("/api/llm/search")
async def ai_search(payload: Dict[str, Any]):
    query = str(payload.get("query", "")).strip()
    if not query:
        raise HTTPException(status_code=400, detail="请输入搜索内容")
    mode = str(payload.get("mode", "")).strip().lower()
    with db_session() as conn:
        row = conn.execute("SELECT * FROM llm_models WHERE is_default = 1 ORDER BY id LIMIT 1").fetchone()
        if not row:
            raise HTTPException(status_code=400, detail="AI 模型服务暂不可用")
        config = dict(row)
        if mode in ("fast", "accurate"):
            fast_id = _get_int_setting(conn, "llm_fast_model_id")
            accurate_id = _get_int_setting(conn, "llm_accurate_model_id")
            selected_id = fast_id if mode == "fast" else accurate_id
            if selected_id:
                selected_row = conn.execute(
                    "SELECT * FROM llm_models WHERE id = ?",
                    (selected_id,),
                ).fetchone()
                if selected_row:
                    config = dict(selected_row)
        devices = _fetch_devices(conn)
    ai_candidates = _filter_ai_devices(devices)
    device_payload = [
        {
            "id": item["id"],
            "model": item["model"],
            "type": item["type"],
            "vendor": item["vendor_name"],
            "system": item["system_name"],
            "version": item["system_version"],
            "status": item["status"],
            "resolution": item["resolution"],
            "arch": item["arch"],
            "cpu": item["cpu"],
            "boot_password": item["boot_password"],
            "notes": item["notes"],
            "loan_status": item["loan_status"],
            "borrower_name": item["borrower_name"],
            "borrowed_at": item["borrowed_at"],
            "expected_return_at": item["expected_return_at"],
        }
        for item in ai_candidates
    ]
    system_prompt = (
        "你是设备匹配助手。根据用户需求，从设备列表中挑选最合适的设备 ID。"
        "仅返回 JSON，不要输出多余文字或代码块。"
        "返回 JSON 格式：{\"device_ids\":[1,2],\"reason\":\"...\"}"
    )
    user_prompt = f"用户需求: {query}\n设备列表: {json.dumps(device_payload, ensure_ascii=False)}"
    client: httpx.AsyncClient = app.state.http_client
    try:
        result = await call_llm(client, config, system_prompt, user_prompt)
    except LLMError as exc:
        raise HTTPException(status_code=400, detail=f"AI 模型服务暂不可用: {exc}")
    device_ids = result.get("device_ids") or []
    if not isinstance(device_ids, list):
        device_ids = []
    matched = [item for item in ai_candidates if item["id"] in device_ids]
    return {"items": matched, "ai_reason": result.get("reason", "")}
